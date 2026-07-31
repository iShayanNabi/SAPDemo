"""Integration tests over the generated spend dataset.

The scenario manifest states what each injected condition should produce. These
tests read that manifest back and check the engine actually produces it, which
is what makes the sample data a regression test rather than decoration.

They also run the full HTTP journey for CSV, XLSX and JSON, because the three
files carry three different header conventions and must all reach the same
canonical result.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modules.spend.analytics import build_analytics
from app.modules.spend.field_definitions import REGISTRY
from app.modules.spend.metrics import calculate_metrics, calculate_supplier_spend
from app.modules.spend.normalizer import normalize_spend_dataframe
from app.modules.spend.savings import calculate_savings
from app.services.files.readers import read_tabular
from app.services.tabular.mapping import suggest_mapping

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def spend_analysis(request, spend_config):
    """Run the full pipeline once over the sample CSV and reuse the result."""
    path = Path(request.getfixturevalue("spend_sample_csv_path"))
    read_result = read_tabular(path.read_bytes(), ".csv")
    mapping = suggest_mapping(read_result.source_columns, REGISTRY)
    dataset = normalize_spend_dataframe(read_result.dataframe, mapping.mapping, spend_config)

    supplier_rows = calculate_supplier_spend(dataset.frame, spend_config)
    savings = calculate_savings(dataset.frame, spend_config, supplier_rows)
    metrics = calculate_metrics(
        dataset.frame, spend_config, supplier_rows, estimated_savings=savings.total_estimated_saving
    )
    analytics = build_analytics(dataset.frame, spend_config, supplier_rows)
    return dataset, metrics, analytics, savings, supplier_rows, mapping


def scenario(manifest: dict, scenario_id: str) -> dict:
    """Look up one documented scenario."""
    for entry in manifest["scenarios"]:
        if entry["scenario_id"] == scenario_id:
            return entry
    raise AssertionError(f"Scenario {scenario_id} is not documented in the manifest")


# ---------------------------------------------------------------------------
# Dataset shape
# ---------------------------------------------------------------------------
def test_dataset_covers_at_least_24_months(spend_analysis, spend_scenario_manifest):
    _dataset, _metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    assert spend_scenario_manifest["months_covered"] >= 24
    assert len(analytics["monthly_spend"]) >= 24


def test_dataset_size_and_variety(spend_analysis):
    dataset, metrics, _analytics, _savings, _suppliers, _mapping = spend_analysis
    frame = dataset.frame

    assert len(frame) >= 1000
    assert metrics.supplier_count >= 50
    assert frame["category"].nunique() >= 5
    assert frame["currency"].nunique() >= 3


def test_headers_map_without_manual_correction(spend_analysis):
    _dataset, _metrics, _analytics, _savings, _suppliers, mapping = spend_analysis
    assert mapping.unmapped_columns == []
    assert mapping.mapping["EBELN"] == "po_number"
    assert mapping.mapping["POSTING_DATE"] == "transaction_date"
    assert mapping.mapping["SPEND_CATEGORY"] == "category"
    assert mapping.mapping["BASELINE_PRICE"] == "baseline_price"


def test_sample_data_has_no_type_errors(spend_analysis):
    dataset, _metrics, _analytics, _savings, _suppliers, _mapping = spend_analysis
    failures = [i for i in dataset.issues if i.issue_type == "type_conversion_failed"]
    assert failures == []


def test_headline_metrics_are_coherent(spend_analysis):
    _dataset, metrics, _analytics, _savings, _suppliers, _mapping = spend_analysis

    assert metrics.total_spend > 0
    assert metrics.contracted_spend + metrics.non_contracted_spend == pytest.approx(
        metrics.total_spend, rel=1e-6
    )
    assert 0 <= metrics.maverick_spend_pct <= 100
    assert 0 <= metrics.spend_under_management_pct <= 100
    assert metrics.maverick_spend <= metrics.non_contracted_spend
    assert metrics.top_five_supplier_share_pct >= metrics.top_supplier_share_pct
    assert 0 <= metrics.supplier_concentration_hhi <= 10000


# ---------------------------------------------------------------------------
# Documented scenarios
# ---------------------------------------------------------------------------
def test_sc01_maverick_spend_is_detected(spend_analysis, spend_scenario_manifest):
    _dataset, metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-01")

    assert metrics.maverick_spend > 0
    flagged = {row["value"] for row in analytics["maverick_spend"]}
    missing = [s for s in documented["suppliers"] if s not in flagged]
    assert not missing, f"maverick suppliers not detected: {missing}"


def test_sc02_contract_leakage_is_detected(spend_analysis, spend_scenario_manifest):
    _dataset, _metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-02")

    leaking = {row["value"]: row for row in analytics["contract_leakage"]}
    for supplier_id in documented["suppliers"]:
        assert supplier_id in leaking, f"leakage not detected for {supplier_id}"
        assert leaking[supplier_id]["leaked_spend_base"] > 0
        assert leaking[supplier_id]["contracted_spend_base"] > 0


def test_sc03_supplier_concentration_is_detected(spend_analysis, spend_scenario_manifest):
    _dataset, _metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-03")

    flagged = [row for row in analytics["supplier_concentration"] if row["exceeds_warning_threshold"]]
    assert flagged, "no concentrated material group was reported"
    concentrated_suppliers = {row["top_supplier_id"] for row in flagged}
    assert set(documented["suppliers"]).issubset(concentrated_suppliers)
    assert any(row["top_supplier_share_pct"] >= 99.0 for row in flagged)


def test_sc04_price_variance_is_detected(spend_analysis, spend_scenario_manifest):
    _dataset, metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-04")

    assert metrics.price_variance_base > 0
    assert metrics.price_variance_line_count > 0
    flagged = {row["value"] for row in analytics["purchase_price_variance"]
               if row["exceeds_alert_threshold"]}
    missing = [m for m in documented["materials"] if m not in flagged]
    assert not missing, f"price variance not detected for {missing}"


def test_sc05_tail_spend_is_detected(spend_analysis, spend_scenario_manifest):
    _dataset, metrics, analytics, _savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-05")

    assert metrics.tail_supplier_count >= 20
    assert metrics.tail_spend > 0
    tail = {row["supplier_id"] for row in analytics["tail_spend_suppliers"]}
    missing = [s for s in documented["suppliers"] if s not in tail]
    assert not missing, f"documented tail suppliers not classified as tail: {missing}"


def test_sc06_preferred_supplier_gap_produces_an_opportunity(
    spend_analysis, spend_scenario_manifest
):
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-06")

    migrations = {o.scope_value for o in savings.opportunities if o.rule_id == "SAV-05"}
    assert migrations, "SAV-05 produced no opportunity"
    assert migrations.intersection(documented["materials"]), (
        "no documented preferred-supplier material was proposed for migration"
    )


def test_sc07_fragmented_group_produces_a_consolidation_opportunity(
    spend_analysis, spend_scenario_manifest
):
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    scenario(spend_scenario_manifest, "SC-07")

    consolidations = [o for o in savings.opportunities if o.rule_id == "SAV-03"]
    assert consolidations, "SAV-03 produced no opportunity"
    assert any(o.evidence["supplier_count"] >= 4 for o in consolidations)


def test_sc08_price_dispersion_produces_a_harmonisation_opportunity(
    spend_analysis, spend_scenario_manifest
):
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    documented = scenario(spend_scenario_manifest, "SC-08")

    harmonisations = {o.scope_value for o in savings.opportunities if o.rule_id == "SAV-01"}
    assert harmonisations, "SAV-01 produced no opportunity"
    assert harmonisations.intersection(documented["materials"]), (
        "no documented dispersed material was proposed for harmonisation"
    )


def test_every_savings_rule_fires_on_the_sample(spend_analysis, spend_config):
    """A configured model that never fires on the demo data cannot be reviewed."""
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    produced = {o.rule_id for o in savings.opportunities}
    silent = sorted(set(spend_config.savings_rules) - produced)
    assert silent == [], f"savings rules produced nothing on the demo data: {silent}"


def test_no_savings_rule_errored(spend_analysis):
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    assert savings.rule_errors == []


def test_every_opportunity_is_explainable(spend_analysis):
    _dataset, _metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    for opportunity in savings.opportunities:
        assert opportunity.method and len(opportunity.method) > 40
        assert opportunity.evidence
        assert opportunity.is_estimate is True
        assert opportunity.estimated_saving_base <= opportunity.gross_saving_base
        assert opportunity.addressable_spend_base >= 0


def test_savings_stay_a_small_share_of_spend(spend_analysis):
    """A model claiming to save most of the spend would not be credible."""
    _dataset, metrics, _analytics, savings, _suppliers, _mapping = spend_analysis
    share = savings.total_estimated_saving / metrics.total_spend * 100.0
    assert 0 < share < 25.0, f"estimated savings are {share:.1f}% of spend"


def test_results_match_the_recorded_baseline(spend_analysis, spend_sample_csv_path):
    """Guards against an accidental change in the metrics or the generator."""
    baseline_path = spend_sample_csv_path.parent / "expected_spend_baseline.json"
    if not baseline_path.is_file():
        pytest.skip("Baseline not generated; run scripts/generate_spend_sample_data.py")

    baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
    dataset, metrics, analytics, savings, _suppliers, _mapping = spend_analysis

    assert len(dataset.frame) == baseline["record_count"]
    assert savings.by_rule() == baseline["opportunities_by_rule"]
    assert len(savings.opportunities) == baseline["opportunity_count"]

    recorded = baseline["metrics"]
    for key in (
        "total_spend", "purchase_order_count", "supplier_count", "contracted_spend",
        "maverick_spend", "tail_spend", "supplier_concentration_hhi",
        "top_supplier_share_pct", "price_variance_base", "estimated_savings_opportunity",
    ):
        assert getattr(metrics, key) == pytest.approx(recorded[key]), key

    assert {k: len(v) for k, v in analytics.items()} == baseline["analytics_row_counts"]


# ---------------------------------------------------------------------------
# Full HTTP journey
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("fixture_name", "mime"),
    [
        ("spend_sample_csv_path", "text/csv"),
        ("spend_sample_xlsx_path", XLSX_MIME),
        ("spend_sample_json_path", "application/json"),
    ],
)
def test_end_to_end_journey_per_format(api_client, request, fixture_name, mime):
    """Upload -> analyse -> drill down -> opportunities -> export."""
    path: Path = request.getfixturevalue(fixture_name)

    upload = api_client.post(
        "/api/v1/spend/upload", files={"file": (path.name, path.read_bytes(), mime)}
    ).json()["data"]
    assert upload["is_analyzable"], upload["missing_required_fields"]

    analysis = api_client.post(
        "/api/v1/spend/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": True},
    ).json()["data"]

    assert analysis["status"] == "completed"
    assert analysis["metrics"]["total_spend"] > 0
    assert analysis["metrics"]["supplier_count"] >= 50
    assert len(analysis["analytics"]["monthly_spend"]) >= 24
    assert analysis["ai_narrative"]["origin"] == "mock_ai"
    assert analysis["opportunities"]

    supplier = analysis["supplier_spend"][0]
    drill = api_client.get(
        f"/api/v1/spend/analyses/{analysis['analysis_id']}/transactions",
        params={"dimension": "supplier", "value": supplier["supplier_id"], "limit": 1000},
    ).json()["data"]
    assert drill["total"] == supplier["transaction_count"]
    assert drill["total_spend_base"] == pytest.approx(supplier["spend_base"])

    export = api_client.get(
        f"/api/v1/spend/analyses/{analysis['analysis_id']}/export", params={"format": "xlsx"}
    )
    workbook = load_workbook(io.BytesIO(export.content))
    assert "Savings Opportunities" in workbook.sheetnames


def test_all_three_formats_produce_the_same_metrics(
    api_client, spend_sample_csv_path, spend_sample_xlsx_path, spend_sample_json_path
):
    """The three files carry the same data under different headers."""
    results = []
    for path, mime in (
        (spend_sample_csv_path, "text/csv"),
        (spend_sample_xlsx_path, XLSX_MIME),
        (spend_sample_json_path, "application/json"),
    ):
        upload = api_client.post(
            "/api/v1/spend/upload", files={"file": (path.name, path.read_bytes(), mime)}
        ).json()["data"]
        analysis = api_client.post(
            "/api/v1/spend/analyze",
            json={"upload_id": upload["upload_id"], "generate_ai_summary": False},
        ).json()["data"]
        metrics = analysis["metrics"]
        results.append(
            (
                round(metrics["total_spend"], 2),
                metrics["supplier_count"],
                metrics["line_item_count"],
                round(metrics["maverick_spend"], 2),
                len(analysis["opportunities"]),
            )
        )

    assert results[0] == results[1] == results[2]


def test_filtered_analysis_is_internally_consistent(api_client, spend_sample_csv_path):
    """KPIs, breakdowns and drill-down must all describe the same slice."""
    upload = api_client.post(
        "/api/v1/spend/upload",
        files={"file": (spend_sample_csv_path.name, spend_sample_csv_path.read_bytes(), "text/csv")},
    ).json()["data"]

    full = api_client.post(
        "/api/v1/spend/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": False},
    ).json()["data"]

    filtered = api_client.post(
        "/api/v1/spend/analyze",
        json={
            "upload_id": upload["upload_id"],
            "filters": {"date_from": "2025-01-01", "date_to": "2025-12-31"},
            "generate_ai_summary": False,
        },
    ).json()["data"]

    assert filtered["filtered_record_count"] < full["filtered_record_count"]
    assert filtered["metrics"]["total_spend"] < full["metrics"]["total_spend"]
    assert len(filtered["analytics"]["monthly_spend"]) == 12

    # The breakdown total must equal the headline metric.
    category_total = sum(
        row["spend_base"] for row in filtered["analytics"]["spend_by_category"]
    )
    assert category_total == pytest.approx(filtered["metrics"]["total_spend"], rel=1e-4)

    # And the drill-down must reconcile with the breakdown.
    category = filtered["analytics"]["spend_by_category"][0]
    drill = api_client.get(
        f"/api/v1/spend/analyses/{filtered['analysis_id']}/transactions",
        params={"dimension": "category", "value": category["value"], "limit": 1000},
    ).json()["data"]
    assert drill["total"] == category["transaction_count"]


def test_supplier_breakdown_sums_to_total_spend(api_client, spend_sample_csv_path):
    upload = api_client.post(
        "/api/v1/spend/upload",
        files={"file": (spend_sample_csv_path.name, spend_sample_csv_path.read_bytes(), "text/csv")},
    ).json()["data"]
    analysis = api_client.post(
        "/api/v1/spend/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": False},
    ).json()["data"]

    supplier_total = sum(row["spend_base"] for row in analysis["supplier_spend"])
    assert supplier_total == pytest.approx(analysis["metrics"]["total_spend"], rel=1e-4)

    monthly_total = sum(row["spend_base"] for row in analysis["analytics"]["monthly_spend"])
    assert monthly_total == pytest.approx(analysis["metrics"]["total_spend"], rel=1e-4)


def test_po_risk_sample_file_also_works_here(api_client, sample_csv_path):
    """A purchase order extract should analyse without any spend-specific columns.

    This is the payoff of reusing module 1's field contract: contract status is
    derived from the contract number, and the order date stands in for the
    transaction date.
    """
    upload = api_client.post(
        "/api/v1/spend/upload",
        files={"file": (sample_csv_path.name, sample_csv_path.read_bytes(), "text/csv")},
    ).json()["data"]
    assert upload["is_analyzable"], upload["missing_required_fields"]

    analysis = api_client.post(
        "/api/v1/spend/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": False},
    ).json()["data"]

    assert analysis["metrics"]["total_spend"] > 0
    assert analysis["metrics"]["contracted_spend"] > 0
    assert analysis["analytics"]["monthly_spend"]
