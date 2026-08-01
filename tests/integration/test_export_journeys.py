"""The two exports added in Phase 5, driven over HTTP against the bundled data.

These are the journeys the acceptance criterion is about: load the demo file,
run the module, download the report, and check the file says what the API said.
They run against the committed sample datasets rather than hand-built rows, so a
sample file that stops matching its module fails here rather than in front of
somebody following the README.

The assertions worth their space are the ones comparing the *file* to the *API
response*. A report is a second serialisation of the same data, and a second
serialisation is a second chance to disagree.
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

CSV_MIME = "text/csv"
V1 = "/api/v1"


def _ok(response) -> dict:
    """Unwrap the shared envelope from any 2xx.

    The modules disagree about which 2xx they use - starting an interview
    session is a 201, calculating a risk assessment is a 200 - and that is a
    route-level decision this file has no business asserting on.
    """
    assert 200 <= response.status_code < 300, (
        f"{response.request.method} {response.request.url} -> "
        f"{response.status_code}: {response.text[:500]}"
    )
    body = response.json()
    assert body["success"] is True
    return body["data"]


def _sheet_text(content: bytes, name: str) -> str:
    workbook = load_workbook(io.BytesIO(content))
    return "\n".join(
        " | ".join("" if cell.value is None else str(cell.value) for cell in row)
        for row in workbook[name].iter_rows()
    )


# ---------------------------------------------------------------------------
# Module 5
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def risk_assessment(api_client, supplier_risk_sample_csv_path, supplier_risk_event_sample_csv_path):
    """The demo portfolio, scored, over HTTP."""
    profiles = _ok(
        api_client.post(
            f"{V1}/supplier-risk/upload",
            data={"dataset": "profiles"},
            files={
                "file": (
                    supplier_risk_sample_csv_path.name,
                    supplier_risk_sample_csv_path.read_bytes(),
                    CSV_MIME,
                )
            },
        )
    )
    _ok(
        api_client.post(
            f"{V1}/supplier-risk/upload",
            data={"dataset": "events", "dataset_id": profiles["dataset_id"]},
            files={
                "file": (
                    supplier_risk_event_sample_csv_path.name,
                    supplier_risk_event_sample_csv_path.read_bytes(),
                    CSV_MIME,
                )
            },
        )
    )
    return _ok(
        api_client.post(
            f"{V1}/supplier-risk/calculate",
            json={"dataset_id": profiles["dataset_id"], "generate_ai_summary": True},
        )
    )


class TestSupplierRiskExportJourney:
    def test_the_whole_portfolio_downloads_in_every_format(self, api_client, risk_assessment):
        assessment_id = risk_assessment["assessment_id"]
        for export_format, extension in (("xlsx", ".xlsx"), ("csv", ".csv"), ("json", ".json")):
            response = api_client.get(
                f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                params={"format": export_format},
            )
            assert response.status_code == 200, response.text
            assert extension in response.headers["content-disposition"]
            assert len(response.content) > 500, f"the {export_format} export looks empty"

    def test_the_file_reports_the_same_suppliers_the_api_did(self, api_client, risk_assessment):
        """A report is a second serialisation, and a second chance to disagree."""
        assessment_id = risk_assessment["assessment_id"]

        document = json.loads(
            api_client.get(
                f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                params={"format": "json"},
            ).content
        )

        api_ids = [supplier["supplier_id"] for supplier in risk_assessment["suppliers"]]
        file_ids = [supplier["supplier_id"] for supplier in document["suppliers"]]
        assert file_ids == api_ids

        api_scores = {s["supplier_id"]: s["overall_score"] for s in risk_assessment["suppliers"]}
        for supplier in document["suppliers"]:
            assert supplier["overall_score"] == api_scores[supplier["supplier_id"]]

    def test_every_exported_score_can_be_rebuilt_from_its_contributions(
        self, api_client, risk_assessment
    ):
        """Across the whole demo portfolio, not just one supplier."""
        document = json.loads(
            api_client.get(
                f"{V1}/supplier-risk/assessments/{risk_assessment['assessment_id']}/export",
                params={"format": "json"},
            ).content
        )
        assert len(document["profiles"]) > 10, "the demo portfolio is too small to be a check"

        rebuilt = 0
        for profile in document["profiles"]:
            contributions = sum(
                category["contribution"] or 0.0 for category in profile["categories"]
            )
            if profile["overall_score"] is None:
                # SRK-09: too few categories could be scored, so the engine
                # withholds the score rather than publishing a partial one. The
                # report has to withhold it too - and must not quietly print the
                # contributions it does have as though they were the total.
                assert profile["limited_data"] is True, (
                    f"{profile['supplier_id']}: no overall score but not flagged limited_data"
                )
                assert profile["unscored_categories"]
                continue
            assert contributions == pytest.approx(profile["overall_score"], abs=0.05), (
                f"{profile['supplier_id']}: contributions do not rebuild the overall score"
            )
            rebuilt += 1
        assert rebuilt > 10, "almost every demo supplier is scored; the check did nothing"

    def test_the_missing_data_supplier_is_exported_as_unscored_not_as_zero(
        self, api_client, risk_assessment, supplier_risk_scenario_manifest
    ):
        """The demo dataset has a deliberate hole in it. The report has to show it.

        A supplier scored on four categories and a supplier scored on ten both
        produce one number. Only the report says which is which.
        """
        document = json.loads(
            api_client.get(
                f"{V1}/supplier-risk/assessments/{risk_assessment['assessment_id']}/export",
                params={"format": "json"},
            ).content
        )

        incomplete = [
            profile for profile in document["profiles"] if profile["unscored_categories"]
        ]
        assert incomplete, "the demo dataset no longer contains a missing-data supplier"

        profile = incomplete[0]
        assert profile["data_completeness_pct"] < 100.0
        unscored = {
            category["category"]
            for category in profile["categories"]
            if not category["data_available"]
        }
        assert unscored, "a category was reported unscored but every category has data"
        for category in profile["categories"]:
            if not category["data_available"]:
                assert category["score"] is None, "an unscored category was given a score"

    def test_the_workbook_carries_the_evidence_and_the_actions(
        self, api_client, risk_assessment
    ):
        content = api_client.get(
            f"{V1}/supplier-risk/assessments/{risk_assessment['assessment_id']}/export",
            params={"format": "xlsx"},
        ).content

        evidence = _sheet_text(content, "Evidence & Missing Data")
        assert "Data Completeness %" in evidence
        assert "Unscored Categories" in evidence

        actions = _sheet_text(content, "Recommended Actions")
        assert "Trigger" in actions
        assert len(actions.splitlines()) > 5, "no recommended action reached the report"

        summary = _sheet_text(content, "Summary")
        assert "NO EXTERNAL DATA SOURCE" in summary

    def test_one_supplier_can_be_exported_on_its_own(self, api_client, risk_assessment):
        supplier_id = risk_assessment["suppliers"][0]["supplier_id"]

        document = json.loads(
            api_client.get(
                f"{V1}/supplier-risk/assessments/{risk_assessment['assessment_id']}/export",
                params={"format": "json", "supplier_id": supplier_id},
            ).content
        )

        assert [p["supplier_id"] for p in document["profiles"]] == [supplier_id]
        assert document["methodology"], "the single-supplier report dropped the methodology"
        assert "NO EXTERNAL DATA SOURCE" in document["disclaimer"]


# ---------------------------------------------------------------------------
# Module 10
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def completed_session(api_client) -> dict:
    """A finished session over the bundled question bank."""
    session = _ok(
        api_client.post(
            f"{V1}/interviews/start",
            json={
                "tracks": ["sap_mm"],
                "mode": "practice",
                "question_count": 3,
                "seed": 20260801,
                "candidate_name": "Export journey",
            },
        )
    )
    answer = (
        "Goods receipt and invoice receipt are matched through the GR/IR clearing account. "
        "The purchase order, the goods receipt and the supplier invoice form the three-way "
        "match, so the company only settles invoices for goods it actually received. An aged "
        "GR/IR balance means one of the three documents is missing."
    )
    for _ in range(3):
        _ok(
            api_client.post(
                f"{V1}/interviews/{session['session_id']}/answer",
                json={"answer_text": answer, "seconds_spent": 90},
            )
        )
    return _ok(api_client.post(f"{V1}/interviews/{session['session_id']}/complete", json={}))


class TestInterviewExportJourney:
    def test_the_session_downloads_in_every_format(self, api_client, completed_session):
        session_id = completed_session["session_id"]
        for export_format, extension in (
            ("xlsx", ".xlsx"),
            ("csv", ".csv"),
            ("json", ".json"),
            ("pdf", ".pdf"),
        ):
            response = api_client.get(
                f"{V1}/interviews/{session_id}/export", params={"format": export_format}
            )
            assert response.status_code == 200, response.text
            assert extension in response.headers["content-disposition"]
            assert len(response.content) > 500, f"the {export_format} export looks empty"

    def test_the_file_reports_the_same_marks_the_api_did(self, api_client, completed_session):
        document = json.loads(
            api_client.get(
                f"{V1}/interviews/{completed_session['session_id']}/export",
                params={"format": "json"},
            ).content
        )

        assert document["summary"]["average_score"] == completed_session["summary"][
            "average_score"
        ]
        api_scores = [
            answer["score"]["overall_score"] for answer in completed_session["answers"]
        ]
        file_scores = [answer["score"]["overall_score"] for answer in document["answers"]]
        assert file_scores == api_scores

    def test_every_question_and_answer_reaches_the_report(
        self, api_client, completed_session
    ):
        document = json.loads(
            api_client.get(
                f"{V1}/interviews/{completed_session['session_id']}/export",
                params={"format": "json"},
            ).content
        )

        assert len(document["answers"]) == completed_session["summary"]["question_count"]
        for answer in document["answers"]:
            assert answer["question"]["question"], "a question text is missing"
            assert answer["answer_text"], "an answer text is missing"
            assert answer["score"]["dimensions"]
            assert answer["score"]["concept_matches"]

    def test_the_transcript_prints_what_credited_each_concept(
        self, api_client, completed_session
    ):
        """Keyword matching has a ceiling; the report has to be arguable with."""
        content = api_client.get(
            f"{V1}/interviews/{completed_session['session_id']}/export",
            params={"format": "pdf"},
        ).content

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "credited by" in text, "no concept shows the keyword that credited it"
        assert "PRACTICE FEEDBACK" in text
        assert len(reader.pages) >= 1 + len(completed_session["answers"])

    def test_a_dimension_that_did_not_apply_is_exported_as_not_applicable(
        self, api_client, completed_session
    ):
        """The bundled bank has questions that test no architecture concept."""
        content = api_client.get(
            f"{V1}/interviews/{completed_session['session_id']}/export",
            params={"format": "xlsx"},
        ).content

        text = _sheet_text(content, "Dimension Scores")
        assert "Applicable" in text
        assert "No" in text, "no dimension was reported as not applicable"

    def test_the_study_plan_reaches_the_report(self, api_client, completed_session):
        content = api_client.get(
            f"{V1}/interviews/{completed_session['session_id']}/export",
            params={"format": "xlsx"},
        ).content

        text = _sheet_text(content, "Study Plan")
        assert "Recommended topics" in text
        assert "Most missed concepts" in text

    def test_the_csv_has_a_row_for_every_answer(self, api_client, completed_session):
        text = api_client.get(
            f"{V1}/interviews/{completed_session['session_id']}/export",
            params={"format": "csv"},
        ).content.decode("utf-8-sig")

        rows = [line for line in text.splitlines() if line.strip()]
        # One header, plus one row per answer - answers wrap across lines when a
        # field contains a newline, so the count is a floor rather than equality.
        assert len(rows) >= 1 + len(completed_session["answers"])
        assert rows[0].startswith("#,Question ID")
