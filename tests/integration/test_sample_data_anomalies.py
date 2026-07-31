"""Integration tests over the generated demo dataset.

The anomaly manifest states, record by record, what the engine is expected to
find. These tests read that manifest and check the engine actually finds it -
which is what makes the sample data a regression test rather than decoration.

They also run the full HTTP journey (upload -> analyse -> findings -> export)
for CSV, XLSX and JSON, because the three files use three different header
conventions and must all end up with the same canonical mapping.
"""

from __future__ import annotations

import csv
import io
import json
from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modules.po_risk.column_mapping import suggest_mapping
from app.modules.po_risk.engine import RiskEngine
from app.modules.po_risk.normalizer import normalize_dataframe
from app.schemas.common import Severity
from app.services.files.readers import read_tabular

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SEVERITY_RANK = {"low": 1, "medium": 2, "high": 3, "critical": 4}


@pytest.fixture(scope="module")
def sample_analysis(request, rule_config):
    """Run the engine once over the sample CSV and reuse the result."""
    path = Path(request.getfixturevalue("sample_csv_path"))
    read_result = read_tabular(path.read_bytes(), ".csv")
    mapping = suggest_mapping(read_result.source_columns)
    dataset = normalize_dataframe(read_result.dataframe, mapping.mapping, rule_config)
    result = RiskEngine(rule_config).run(dataset.frame)
    return dataset, result, mapping


@pytest.fixture(scope="module")
def manifest(request) -> list[dict[str, str]]:
    """The documented anomalies."""
    path = Path(request.getfixturevalue("anomaly_manifest_path"))
    with path.open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


# ---------------------------------------------------------------------------
# Dataset shape
# ---------------------------------------------------------------------------
def test_sample_dataset_meets_the_size_requirements(sample_analysis):
    dataset, _result, _mapping = sample_analysis
    frame = dataset.frame

    assert len(frame) >= 1000, "at least 1,000 line items"
    assert frame["po_number"].nunique() >= 500, "at least 500 purchase orders"
    assert frame["supplier_id"].nunique() >= 50, "at least 50 suppliers"


def test_sap_headers_map_without_manual_correction(sample_analysis):
    _dataset, _result, mapping = sample_analysis
    assert mapping.is_analyzable
    assert mapping.unmapped_columns == []
    assert mapping.mapping["EBELN"] == "po_number"
    assert mapping.mapping["LIFNR"] == "supplier_id"


def test_clean_sample_data_has_no_type_errors(sample_analysis):
    dataset, _result, _mapping = sample_analysis
    failures = [issue for issue in dataset.issues if issue.issue_type == "type_conversion_failed"]
    assert failures == []


# ---------------------------------------------------------------------------
# Manifest verification
# ---------------------------------------------------------------------------
def test_manifest_covers_every_rule(manifest):
    from app.modules.po_risk.rules import RULE_IDS

    documented = {entry["expected_rule_id"] for entry in manifest}
    assert documented == set(RULE_IDS), "every rule needs at least one documented anomaly"


def test_every_rule_produces_findings_on_the_sample(sample_analysis):
    _dataset, result, _mapping = sample_analysis
    silent = [
        execution.rule_id
        for execution in result.executions
        if execution.enabled and execution.findings_count == 0
    ]
    assert silent == [], f"rules produced no finding on the demo data: {silent}"


def test_no_rule_raised_an_error_on_the_sample(sample_analysis):
    _dataset, result, _mapping = sample_analysis
    assert result.rule_errors == []


@pytest.mark.parametrize(
    "rule_id",
    ["PO-R001", "PO-R002", "PO-R003", "PO-R004", "PO-R005", "PO-R006", "PO-R007", "PO-R008",
     "PO-R009", "PO-R010", "PO-R011", "PO-R012", "PO-R013", "PO-R014", "PO-R015", "PO-R016",
     "PO-R017", "PO-R018", "PO-R019", "PO-R020"],
)
def test_documented_anomalies_are_detected(sample_analysis, manifest, rule_id):
    """Every documented record must be found by the rule that was targeted."""
    _dataset, result, _mapping = sample_analysis

    # A finding is keyed by purchase order, except for portfolio level rules
    # such as PO-R018 which report against a supplier and carry no PO number.
    severities: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for finding in result.findings:
        if finding.po_number is not None:
            severities[(finding.rule_id, "purchase_order", finding.po_number)].append(
                finding.severity.value
            )
        if finding.supplier_id is not None:
            severities[(finding.rule_id, "supplier", finding.supplier_id)].append(
                finding.severity.value
            )

    expected = [entry for entry in manifest if entry["expected_rule_id"] == rule_id]
    assert expected, f"no manifest entry for {rule_id}"

    for entry in expected:
        scope = entry["expected_scope"]
        target = entry["supplier_id"] if scope == "supplier" else entry["po_number"]
        actual = severities.get((rule_id, scope, target))
        assert actual, f"{rule_id} did not flag documented {scope} {target}"

        best = max(SEVERITY_RANK[value] for value in actual)
        assert best >= SEVERITY_RANK[entry["expected_severity"]], (
            f"{rule_id} on {scope} {target}: expected at least "
            f"{entry['expected_severity']}, got {actual}"
        )


def test_findings_match_the_recorded_baseline(sample_analysis, sample_csv_path):
    """Guards against an accidental change in either the rules or the data."""
    baseline_path = sample_csv_path.parent / "expected_findings_baseline.json"
    if not baseline_path.is_file():
        pytest.skip("Baseline not generated; run scripts/generate_sample_data.py")

    baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
    _dataset, result, _mapping = sample_analysis
    actual = {execution.rule_id: execution.findings_count for execution in result.executions}

    assert actual == baseline["findings_by_rule"]
    assert result.summary["severity_counts"] == baseline["severity_counts"]


def test_every_finding_is_explainable(sample_analysis):
    """A finding without evidence or an action is not actionable."""
    _dataset, result, _mapping = sample_analysis
    for finding in result.findings:
        assert finding.explanation and len(finding.explanation) > 30
        assert finding.recommended_action
        assert finding.evidence, f"{finding.rule_id} produced no evidence"
        assert 0.0 < finding.confidence_score <= 1.0
        assert finding.output_origin.value == "rule_based"
        assert finding.severity in set(Severity)


def test_supplier_watch_list_orders_are_all_flagged(sample_analysis, rule_config):
    """PO-R020 must cover *every* order of a watch-listed supplier."""
    dataset, result, _mapping = sample_analysis
    frame = dataset.frame
    watch_listed = {entry.supplier_id for entry in rule_config.high_risk_suppliers}

    expected_orders = set(
        frame.loc[frame["supplier_id"].isin(watch_listed), "po_number"].dropna().unique()
    )
    flagged = {f.po_number for f in result.findings if f.rule_id == "PO-R020"}
    assert expected_orders <= flagged


def test_concentrated_material_group_is_flagged(sample_analysis):
    _dataset, result, _mapping = sample_analysis
    concentration = [f for f in result.findings if f.rule_id == "PO-R018"]
    assert concentration
    assert concentration[0].evidence["share_pct"] >= 75.0


# ---------------------------------------------------------------------------
# Full HTTP journey for all three file formats
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("fixture_name", "mime"),
    [
        ("sample_csv_path", "text/csv"),
        ("sample_xlsx_path", XLSX_MIME),
        ("sample_json_path", "application/json"),
    ],
)
def test_end_to_end_journey_per_format(api_client, request, fixture_name, mime):
    """Upload -> analyse -> filter -> export, for each supported file type."""
    path: Path = request.getfixturevalue(fixture_name)

    upload = api_client.post(
        "/api/v1/po-risk/upload", files={"file": (path.name, path.read_bytes(), mime)}
    ).json()["data"]
    assert upload["is_analyzable"], upload["missing_required_fields"]
    assert upload["row_count"] >= 1000

    analysis = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": True,
              "rewrite_findings": True},
    ).json()["data"]

    assert analysis["status"] == "completed"
    assert analysis["findings_count"] > 50
    assert analysis["supplier_count"] >= 50
    assert analysis["ai_narrative"]["origin"] == "mock_ai"

    critical = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings",
        params={"severity": "critical", "limit": 500},
    ).json()["data"]
    assert all(f["severity"] == "critical" for f in critical["findings"])

    export = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/export", params={"format": "xlsx"}
    )
    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook["Findings"].max_row == analysis["findings_count"] + 1


def test_all_three_formats_produce_the_same_findings(
    api_client, sample_csv_path, sample_xlsx_path, sample_json_path
):
    """CSV, XLSX and JSON carry the same data under different headers."""
    counts = []
    for path, mime in (
        (sample_csv_path, "text/csv"),
        (sample_xlsx_path, XLSX_MIME),
        (sample_json_path, "application/json"),
    ):
        upload = api_client.post(
            "/api/v1/po-risk/upload", files={"file": (path.name, path.read_bytes(), mime)}
        ).json()["data"]
        analysis = api_client.post(
            "/api/v1/po-risk/analyze",
            json={"upload_id": upload["upload_id"], "generate_ai_summary": False},
        ).json()["data"]
        counts.append(
            (analysis["findings_count"], analysis["critical_count"], analysis["record_count"])
        )

    assert counts[0] == counts[1] == counts[2]


def test_ai_rewrites_are_labelled_separately(api_client, sample_csv_path):
    """AI text must never overwrite the rule-based explanation."""
    upload = api_client.post(
        "/api/v1/po-risk/upload",
        files={"file": (sample_csv_path.name, sample_csv_path.read_bytes(), "text/csv")},
    ).json()["data"]
    analysis = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": True,
              "rewrite_findings": True},
    ).json()["data"]

    findings = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings", params={"limit": 20}
    ).json()["data"]["findings"]

    rewritten = [f for f in findings if f["ai_explanation"]]
    assert rewritten, "expected the mock provider to rewrite the top findings"
    for finding in rewritten:
        assert finding["output_origin"] == "rule_based"
        assert finding["ai_output_origin"] == "mock_ai"
        assert finding["explanation"] != finding["ai_explanation"]
