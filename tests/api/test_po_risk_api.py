"""API tests for the Purchase Order Risk Checker endpoints.

These run against a real FastAPI test client backed by a temporary SQLite
database, so they exercise routing, validation, the service layer, persistence
and the export builders together.
"""

from __future__ import annotations

import io
import json
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from tests.factories import BASE_DATE, make_row, rows_to_csv, rows_to_json, rows_to_xlsx

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SAP_HEADERS = {
    "po_number": "EBELN", "po_item": "EBELP", "supplier_id": "LIFNR", "supplier_name": "NAME1",
    "material": "MATNR", "material_description": "TXZ01", "material_group": "MATKL",
    "company_code": "BUKRS", "purchasing_org": "EKORG", "purchasing_group": "EKGRP",
    "plant": "WERKS", "quantity": "MENGE", "unit_of_measure": "MEINS", "unit_price": "NETPR",
    "currency": "WAERS", "total_value": "NETWR", "order_date": "BEDAT",
    "requested_delivery_date": "EINDT", "actual_delivery_date": "BUDAT",
    "contract_number": "KONNR", "payment_terms": "ZTERM", "approval_status": "FRGKE",
    "created_by": "ERNAM", "changed_by": "AENAM", "change_count": "CHANGE_COUNT",
}


def risky_rows() -> list[dict]:
    """A small dataset containing several known problems."""
    return [
        # Two identical orders three days apart -> PO-R001
        make_row(po_number="4500000001", quantity=50, unit_price=200.0),
        make_row(po_number="4500000002", quantity=50, unit_price=200.0,
                 order_date=BASE_DATE + timedelta(days=3)),
        # High value, not approved -> PO-R009
        make_row(po_number="4500000003", quantity=100, unit_price=500.0,
                 approval_status="Not Approved"),
        # Late delivery -> PO-R010
        make_row(po_number="4500000004",
                 requested_delivery_date=BASE_DATE + timedelta(days=10),
                 actual_delivery_date=BASE_DATE + timedelta(days=45)),
        # Missing contract on a big line -> PO-R008
        make_row(po_number="4500000005", quantity=200, unit_price=400.0, contract_number=None),
    ]


def upload_file(api_client, content: bytes, filename: str, mime: str) -> dict:
    """Upload helper returning the ``data`` block."""
    response = api_client.post(
        "/api/v1/po-risk/upload", files={"file": (filename, content, mime)}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["success"] is True
    return body["data"]


# ---------------------------------------------------------------------------
# System endpoints
# ---------------------------------------------------------------------------
def test_health_endpoint(api_client):
    body = api_client.get("/api/v1/health").json()
    assert body["success"] is True
    assert body["data"]["database_connected"] is True
    assert body["data"]["ai_provider"] == "mock"


def test_index_lists_the_module(api_client):
    body = api_client.get("/").json()
    assert any(module["id"] == "po_risk" for module in body["modules"])


def test_ai_status_never_exposes_a_key(api_client):
    body = api_client.get("/api/v1/po-risk/ai-status").json()["data"]
    assert body["is_mock"] is True
    serialised = json.dumps(body).lower()
    assert "api_key" not in serialised and "sk-" not in serialised


def test_rule_catalogue_exposes_thresholds(api_client):
    rules = api_client.get("/api/v1/po-risk/rules").json()["data"]
    assert len(rules) == 20
    assert all(rule["rule_id"].startswith("PO-R") for rule in rules)
    assert any(rule["params"] for rule in rules)


def test_field_catalogue_lists_sap_aliases(api_client):
    fields = api_client.get("/api/v1/po-risk/fields").json()["data"]
    by_name = {field["name"]: field for field in fields}
    assert "EBELN" in by_name["po_number"]["aliases"]
    assert "LIFNR" in by_name["supplier_id"]["aliases"]
    assert by_name["quantity"]["required"] is True


# ---------------------------------------------------------------------------
# Upload
# ---------------------------------------------------------------------------
def test_upload_valid_csv(api_client):
    data = upload_file(api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv")

    assert data["row_count"] == 5
    assert data["is_analyzable"] is True
    assert data["suggested_mapping"]["EBELN"] == "po_number"
    assert data["missing_required_fields"] == []
    assert len(data["preview_rows"]) == 5


def test_upload_valid_xlsx(api_client):
    data = upload_file(api_client, rows_to_xlsx(risky_rows()), "orders.xlsx", XLSX_MIME)
    assert data["row_count"] == 5
    assert data["is_analyzable"] is True


def test_upload_valid_json(api_client):
    data = upload_file(api_client, rows_to_json(risky_rows()), "orders.json", "application/json")
    assert data["row_count"] == 5
    assert data["is_analyzable"] is True


@pytest.mark.parametrize(
    ("filename", "content", "mime"),
    [
        ("orders.exe", b"MZ\x90\x00binary", "application/octet-stream"),
        ("orders.txt", b"po_number,quantity\n1,2\n", "text/plain"),
        ("orders.csv", b"", "text/csv"),
        ("orders.xlsx", b"po_number,quantity\n1,2\n", XLSX_MIME),
        ("orders.json", b"{not valid json", "application/json"),
    ],
)
def test_invalid_uploads_are_rejected(api_client, filename, content, mime):
    response = api_client.post(
        "/api/v1/po-risk/upload", files={"file": (filename, content, mime)}
    )
    assert response.status_code in (400, 422)
    body = response.json()
    assert body["success"] is False
    assert body["error"]["message"]
    # The safe message must not leak a server path.
    assert "/home/" not in body["error"]["message"]


def test_upload_with_traversal_filename_is_sanitised(api_client):
    data = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "../../etc/orders.csv", "text/csv"
    )
    assert ".." not in data["original_filename"]
    assert data["original_filename"] == "orders.csv"


def test_upload_reports_unmapped_columns(api_client):
    content = rows_to_csv(risky_rows(), SAP_HEADERS).decode()
    header, *rest = content.splitlines()
    modified = "\n".join([header + ",ZZ_CUSTOM", *[line + ",x" for line in rest]])
    data = upload_file(api_client, modified.encode(), "orders.csv", "text/csv")
    assert "ZZ_CUSTOM" in data["unmapped_columns"]


# ---------------------------------------------------------------------------
# Analyze
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def analysis(api_client) -> dict:
    """Upload the risky dataset once and analyse it."""
    upload = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv"
    )
    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "generate_ai_summary": True},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_analysis_detects_expected_rules(analysis):
    rule_ids = {entry["rule_id"] for entry in analysis["kpis"]["rule_counts"]}
    assert {"PO-R001", "PO-R008", "PO-R009", "PO-R010"} <= rule_ids


def test_analysis_reports_kpis(analysis):
    assert analysis["record_count"] == 5
    assert analysis["purchase_order_count"] == 5
    assert analysis["findings_count"] > 0
    assert analysis["risk_score"] > 0
    assert analysis["status"] == "completed"
    assert analysis["base_currency"] == "EUR"


def test_analysis_includes_labelled_mock_ai_summary(analysis):
    narrative = analysis["ai_narrative"]
    assert narrative["available"] is True
    assert narrative["origin"] == "mock_ai"
    assert narrative["prompt_version"].startswith("po_risk_narrative_v")
    assert narrative["summary"]


def test_analysis_exposes_methodology(analysis):
    methodology = analysis["methodology"]
    assert methodology["risk_determination"] == "deterministic Python rules only"
    assert "not connected to any SAP system" in methodology["data_disclaimer"]
    assert methodology["risk_score_method"]


def test_findings_carry_every_required_attribute(analysis, api_client):
    body = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings"
    ).json()["data"]

    assert body["total"] == analysis["findings_count"]
    for finding in body["findings"]:
        for key in (
            "finding_id", "analysis_id", "po_number", "risk_category", "rule_id", "severity",
            "explanation", "evidence", "recommended_action", "confidence_score",
            "estimated_financial_exposure", "created_at",
        ):
            assert key in finding, key
        assert finding["severity"] in {"low", "medium", "high", "critical"}
        assert finding["output_origin"] == "rule_based"


def test_analysis_with_manual_mapping_override(api_client):
    """A user correction must win over the automatic suggestion."""
    rows = risky_rows()
    headers = dict(SAP_HEADERS)
    headers["material_group"] = "ZZ_CATEGORY"  # not auto-mapped with confidence
    upload = upload_file(api_client, rows_to_csv(rows, headers), "orders.csv", "text/csv")

    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={
            "upload_id": upload["upload_id"],
            "column_mapping_overrides": {"ZZ_CATEGORY": "material_group"},
            "generate_ai_summary": False,
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["applied_mapping"]["ZZ_CATEGORY"] == "material_group"
    assert data["ai_narrative"]["available"] is False


def test_analysis_rejects_an_invalid_mapping(api_client):
    upload = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv"
    )
    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"],
              "column_mapping_overrides": {"EBELN": "not_a_real_field"}},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_analysis_rejects_removing_a_required_field(api_client):
    upload = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv"
    )
    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "column_mapping_overrides": {"MENGE": ""}},
    )
    assert response.status_code == 422
    assert "missing_required_fields" in response.json()["error"]["details"]


def test_analysis_with_unknown_upload_returns_404(api_client):
    response = api_client.post(
        "/api/v1/po-risk/analyze", json={"upload_id": "does-not-exist"}
    )
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "not_found"


def test_analysis_can_run_a_rule_subset(api_client):
    upload = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv"
    )
    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "enabled_rules": ["PO-R009"],
              "generate_ai_summary": False},
    )
    data = response.json()["data"]
    assert {entry["rule_id"] for entry in data["kpis"]["rule_counts"]} == {"PO-R009"}


def test_analysis_rejects_unknown_rule_ids(api_client):
    upload = upload_file(
        api_client, rows_to_csv(risky_rows(), SAP_HEADERS), "orders.csv", "text/csv"
    )
    response = api_client.post(
        "/api/v1/po-risk/analyze",
        json={"upload_id": upload["upload_id"], "enabled_rules": ["PO-R999"]},
    )
    assert response.status_code == 422
    assert "PO-R999" in response.json()["error"]["details"]["unknown_rules"]


def test_analyze_rejects_unknown_body_fields(api_client):
    response = api_client.post(
        "/api/v1/po-risk/analyze", json={"upload_id": "x", "surprise": True}
    )
    assert response.status_code == 422


# ---------------------------------------------------------------------------
# Retrieval
# ---------------------------------------------------------------------------
def test_list_analyses(api_client, analysis):
    body = api_client.get("/api/v1/po-risk/analyses", params={"limit": 5}).json()["data"]
    assert body["total"] >= 1
    assert len(body["analyses"]) <= 5
    assert all("analysis_id" in item for item in body["analyses"])


def test_get_single_analysis(api_client, analysis):
    body = api_client.get(f"/api/v1/po-risk/analyses/{analysis['analysis_id']}").json()["data"]
    assert body["analysis_id"] == analysis["analysis_id"]
    assert body["findings_count"] == analysis["findings_count"]


def test_get_unknown_analysis_returns_404(api_client):
    response = api_client.get("/api/v1/po-risk/analyses/unknown-id")
    assert response.status_code == 404


@pytest.mark.parametrize("severity", ["critical", "high", "medium", "low"])
def test_findings_can_be_filtered_by_severity(api_client, analysis, severity):
    body = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings",
        params={"severity": severity},
    ).json()["data"]
    assert all(finding["severity"] == severity for finding in body["findings"])


def test_findings_can_be_filtered_by_rule_and_po(api_client, analysis):
    body = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings",
        params={"rule_id": "PO-R009"},
    ).json()["data"]
    assert body["total"] >= 1
    assert all(finding["rule_id"] == "PO-R009" for finding in body["findings"])
    assert all(finding["po_number"] == "4500000003" for finding in body["findings"])


def test_findings_pagination(api_client, analysis):
    first = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings",
        params={"limit": 1, "offset": 0},
    ).json()["data"]
    second = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/findings",
        params={"limit": 1, "offset": 1},
    ).json()["data"]
    assert len(first["findings"]) == 1
    assert first["findings"][0]["finding_id"] != second["findings"][0]["finding_id"]


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
def test_export_xlsx(api_client, analysis):
    response = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/export", params={"format": "xlsx"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert "attachment;" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content))
    assert set(workbook.sheetnames) == {
        "Summary", "Findings", "Supplier Risk", "Methodology", "Data Quality"
    }
    findings_sheet = workbook["Findings"]
    assert findings_sheet.max_row == analysis["findings_count"] + 1
    headers = [cell.value for cell in findings_sheet[1]]
    assert "Explanation (rule-based)" in headers
    assert "Supporting Evidence" in headers


def test_export_csv(api_client, analysis):
    response = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/export", params={"format": "csv"}
    )
    assert response.status_code == 200
    text = response.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Finding ID,Analysis ID")
    assert len(text.strip().splitlines()) == analysis["findings_count"] + 1


def test_export_json(api_client, analysis):
    response = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/export", params={"format": "json"}
    )
    payload = json.loads(response.content)
    assert payload["report_type"] == "sap_po_risk_analysis"
    assert len(payload["findings"]) == analysis["findings_count"]
    assert len(payload["rule_catalogue"]) == 20
    assert "not connected to a live SAP system" in payload["disclaimer"]


def test_export_rejects_an_unknown_format(api_client, analysis):
    response = api_client.get(
        f"/api/v1/po-risk/analyses/{analysis['analysis_id']}/export", params={"format": "pdf"}
    )
    assert response.status_code == 422


def test_export_of_unknown_analysis_returns_404(api_client):
    response = api_client.get("/api/v1/po-risk/analyses/nope/export")
    assert response.status_code == 404


# ---------------------------------------------------------------------------
# Sample data endpoints
# ---------------------------------------------------------------------------
def test_sample_info(api_client, sample_csv_path):
    body = api_client.get("/api/v1/po-risk/sample/info").json()["data"]
    assert body["available"] is True
    assert body["row_count"] > 1000
    assert body["data_origin"] == "demo_data"


@pytest.mark.parametrize("file_format", ["csv", "xlsx", "json"])
def test_sample_download(api_client, sample_csv_path, file_format):
    response = api_client.get("/api/v1/po-risk/sample", params={"format": file_format})
    assert response.status_code == 200
    assert len(response.content) > 1000
