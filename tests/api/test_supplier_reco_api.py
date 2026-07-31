"""API tests for the Supplier Recommendation Engine endpoints.

Run against a real FastAPI test client backed by a temporary SQLite database, so
routing, validation, the service layer, persistence and the export builders are
exercised together.
"""

from __future__ import annotations

import io
import json

from openpyxl import load_workbook

from tests.factories import supplier_row, supplier_rows_to_csv, supplier_rows_to_xlsx

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TECHNICAL_HEADERS = {
    "supplier_id": "LIFNR", "supplier_name": "NAME1",
    "materials_supplied": "MATERIALS_SUPPLIED", "plants_served": "PLANTS_SERVED",
    "regions_served": "REGIONS_SERVED", "unit_price": "NETPR", "currency": "WAERS",
    "lead_time_days": "PLIFZ", "available_capacity": "AVAILABLE_CAPACITY",
    "on_time_delivery_rate": "OTD", "quality_score": "QUALITY_SCORE",
    "defect_rate": "DEFECT_RATE", "risk_score": "RISK_SCORE", "esg_score": "ESG_SCORE",
    "contract_status": "CONTRACT_STATUS", "contract_expiration": "CONTRACT_EXPIRATION",
    "payment_terms": "ZTERM", "historical_order_count": "ORDER_COUNT",
    "historical_spend": "HISTORICAL_SPEND",
}


def supplier_dataset() -> list[dict]:
    """A small catalogue with a known best/cheapest/ineligible supplier."""
    return [
        supplier_row(
            supplier_id="0000300001", supplier_name="Alpha GmbH",
            materials_supplied=["MAT-1000", "MAT-1010"], plants_served=["1010", "1020"],
            regions_served=["EU"], unit_price=100.0, currency="EUR", lead_time_days=10,
            available_capacity=1000.0, on_time_delivery_rate=98.0, quality_score=92.0,
            defect_rate=1.0, risk_score=20.0, esg_score=85.0, contract_status="Active",
            contract_expiration="2028-01-01", historical_order_count=40, historical_spend=400000.0,
        ),
        supplier_row(
            supplier_id="0000300002", supplier_name="Beta Ltd",
            materials_supplied=["MAT-1000"], plants_served=["1010"], regions_served=["EU"],
            unit_price=70.0, currency="EUR", lead_time_days=25, available_capacity=500.0,
            on_time_delivery_rate=88.0, quality_score=78.0, defect_rate=3.0, risk_score=45.0,
            esg_score=65.0, contract_status="No contract", contract_expiration=None,
            historical_order_count=8, historical_spend=60000.0,
        ),
        supplier_row(
            supplier_id="0000300003", supplier_name="Gamma SA",
            materials_supplied=["MAT-1000"], plants_served=["1010"], regions_served=["EU"],
            unit_price=130.0, currency="EUR", lead_time_days=6, available_capacity=2000.0,
            on_time_delivery_rate=99.0, quality_score=96.0, defect_rate=0.5, risk_score=12.0,
            esg_score=90.0, contract_status="Active", contract_expiration="2029-01-01",
            historical_order_count=60, historical_spend=800000.0,
        ),
        supplier_row(  # high risk -> ineligible at medium tolerance
            supplier_id="0000300004", supplier_name="Delta AG",
            materials_supplied=["MAT-1000"], plants_served=["1010"], regions_served=["EU"],
            unit_price=90.0, currency="EUR", lead_time_days=30, available_capacity=800.0,
            on_time_delivery_rate=70.0, quality_score=60.0, defect_rate=8.0, risk_score=85.0,
            esg_score=40.0, contract_status="No contract", contract_expiration=None,
            historical_order_count=2, historical_spend=5000.0,
        ),
        supplier_row(  # wrong material -> ineligible
            supplier_id="0000300005", supplier_name="Epsilon Oy",
            materials_supplied=["MAT-2000"], plants_served=["2010"], regions_served=["NA"],
            unit_price=110.0, currency="USD", lead_time_days=15, available_capacity=200.0,
            on_time_delivery_rate=95.0, quality_score=88.0, defect_rate=2.0, risk_score=30.0,
            esg_score=78.0, contract_status="Active", contract_expiration="2027-03-01",
            historical_order_count=25, historical_spend=250000.0,
        ),
    ]


REQUIREMENT = {
    "material": "MAT-1000", "quantity": 100, "plant": "1010", "preferred_region": "EU",
    "required_delivery_date": "2026-10-01", "order_date": "2026-08-01",
    "target_price": 120, "currency": "EUR", "risk_tolerance": "medium",
}


def _upload(api_client, content: bytes, filename: str, mime: str) -> dict:
    response = api_client.post("/api/v1/suppliers/upload", files={"file": (filename, content, mime)})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["success"] is True
    return body["data"]


def _load_catalog(api_client) -> str:
    content = supplier_rows_to_csv(supplier_dataset(), headers=TECHNICAL_HEADERS)
    return _upload(api_client, content, "suppliers.csv", "text/csv")["catalog_id"]


# ---------------------------------------------------------------------------
# Catalogues / documentation
# ---------------------------------------------------------------------------
def test_index_lists_the_supplier_module(api_client):
    body = api_client.get("/").json()
    modules = {module["id"]: module for module in body["modules"]}
    assert modules["supplier_recommendation"]["status"] == "available"


def test_field_catalogue_flags_list_fields_and_reuses_aliases(api_client):
    fields = api_client.get("/api/v1/suppliers/fields").json()["data"]
    by_name = {field["name"]: field for field in fields}
    assert by_name["materials_supplied"]["is_list_field"] is True
    assert by_name["plants_served"]["is_list_field"] is True
    assert by_name["unit_price"]["is_list_field"] is False
    # Established SAP alias conventions are reused.
    assert "LIFNR" in by_name["supplier_id"]["aliases"]
    assert "WAERS" in by_name["currency"]["aliases"]


def test_scoring_endpoint_documents_nine_dimensions(api_client):
    data = api_client.get("/api/v1/supplier-recommendations/scoring").json()["data"]
    assert len(data["dimensions"]) == 9
    assert data["weight_total"] == 100.0
    assert round(sum(data["default_weights"].values()), 2) == 100.0
    assert data["eligibility_filters"]


def test_ai_status_never_exposes_a_key(api_client):
    body = api_client.get("/api/v1/supplier-recommendations/ai-status").json()["data"]
    serialised = json.dumps(body).lower()
    assert "api_key" not in serialised and "sk-" not in serialised


# ---------------------------------------------------------------------------
# Upload / suppliers
# ---------------------------------------------------------------------------
def test_upload_loads_suppliers_into_a_catalogue(api_client):
    data = _upload(
        api_client, supplier_rows_to_csv(supplier_dataset(), headers=TECHNICAL_HEADERS),
        "suppliers.csv", "text/csv",
    )
    assert data["supplier_count"] == 5
    assert data["missing_required_fields"] == []
    assert data["catalog_id"]


def test_upload_xlsx_business_labels(api_client):
    content = supplier_rows_to_xlsx(supplier_dataset())
    data = _upload(api_client, content, "suppliers.xlsx", XLSX_MIME)
    assert data["supplier_count"] == 5


def test_list_and_get_suppliers(api_client):
    catalog_id = _load_catalog(api_client)
    listing = api_client.get(
        "/api/v1/suppliers", params={"catalog_id": catalog_id, "material": "MAT-1000"}
    ).json()["data"]
    assert listing["total"] == 4  # four suppliers list MAT-1000
    supplier = api_client.get(
        "/api/v1/suppliers/0000300003", params={"catalog_id": catalog_id}
    ).json()["data"]
    assert supplier["supplier_name"] == "Gamma SA"
    assert "MAT-1000" in supplier["materials_supplied"]


def test_unknown_supplier_returns_404(api_client):
    catalog_id = _load_catalog(api_client)
    response = api_client.get(
        "/api/v1/suppliers/does-not-exist", params={"catalog_id": catalog_id}
    )
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "not_found"


# ---------------------------------------------------------------------------
# Recommend
# ---------------------------------------------------------------------------
def test_recommend_ranks_eligible_and_excludes_ineligible(api_client):
    catalog_id = _load_catalog(api_client)
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "generate_ai_summary": True}
    response = api_client.post("/api/v1/supplier-recommendations/recommend", json=body)
    assert response.status_code == 200, response.text
    data = response.json()["data"]

    # Delta (high risk) and Epsilon (wrong material) are ineligible; three remain.
    assert data["eligible_count"] == 3
    assert data["ineligible_count"] == 2
    by_id = {e["supplier_id"]: e for e in data["results"]}
    assert by_id["0000300004"]["eligibility_status"] == "ineligible"
    assert by_id["0000300005"]["eligibility_status"] == "ineligible"
    # Eligible suppliers carry a rank; ineligible ones do not.
    assert by_id["0000300001"]["rank"] is not None
    assert by_id["0000300004"]["rank"] is None
    # The AI summary is present and labelled mock.
    assert data["ai_narrative"]["origin"] == "mock_ai"


def test_recommend_cost_and_delivery_estimates(api_client):
    catalog_id = _load_catalog(api_client)
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "generate_ai_summary": False}
    data = api_client.post("/api/v1/supplier-recommendations/recommend", json=body).json()["data"]
    alpha = {e["supplier_id"]: e for e in data["results"]}["0000300001"]
    # 100 EUR unit price x 100 quantity.
    assert alpha["estimated_total_cost_base"] == 10000.0
    # order 2026-08-01 + 10 days lead time.
    assert alpha["estimated_delivery_date"] == "2026-08-11"


def test_recommend_with_all_cost_weight_prefers_cheapest(api_client):
    catalog_id = _load_catalog(api_client)
    weights = {
        "cost": 100, "delivery": 0, "quality": 0, "capacity": 0, "risk": 0,
        "esg": 0, "contract": 0, "geographic": 0, "past_performance": 0,
    }
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "weights": weights,
            "generate_ai_summary": False}
    data = api_client.post("/api/v1/supplier-recommendations/recommend", json=body).json()["data"]
    winner = next(e for e in data["results"] if e["rank"] == 1)
    assert winner["supplier_id"] == "0000300002"  # Beta at 70 EUR is cheapest eligible


def test_recommend_rejects_weights_not_summing_to_100(api_client):
    catalog_id = _load_catalog(api_client)
    weights = {
        "cost": 50, "delivery": 10, "quality": 10, "capacity": 5, "risk": 5,
        "esg": 5, "contract": 5, "geographic": 5, "past_performance": 10,  # 105
    }
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "weights": weights}
    response = api_client.post("/api/v1/supplier-recommendations/recommend", json=body)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_recommend_rejects_unknown_body_field(api_client):
    catalog_id = _load_catalog(api_client)
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "surprise": True}
    response = api_client.post("/api/v1/supplier-recommendations/recommend", json=body)
    assert response.status_code == 422


def test_get_recommendation_round_trips(api_client):
    catalog_id = _load_catalog(api_client)
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "generate_ai_summary": False}
    created = api_client.post("/api/v1/supplier-recommendations/recommend", json=body).json()["data"]
    rec_id = created["recommendation_id"]
    fetched = api_client.get(f"/api/v1/supplier-recommendations/{rec_id}").json()["data"]
    assert fetched["recommendation_id"] == rec_id
    assert fetched["eligible_count"] == created["eligible_count"]
    assert len(fetched["results"]) == len(created["results"])


def test_unknown_recommendation_returns_404(api_client):
    response = api_client.get("/api/v1/supplier-recommendations/does-not-exist")
    assert response.status_code == 404


def test_recommend_without_catalogue_after_none_uploaded(api_client):
    # Referencing a non-existent catalogue id is a 404.
    body = {"catalog_id": "no-such-catalog", "requirement": REQUIREMENT}
    response = api_client.post("/api/v1/supplier-recommendations/recommend", json=body)
    assert response.status_code == 404


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
def test_exports_in_all_three_formats(api_client):
    catalog_id = _load_catalog(api_client)
    body = {"catalog_id": catalog_id, "requirement": REQUIREMENT, "generate_ai_summary": False}
    rec_id = api_client.post(
        "/api/v1/supplier-recommendations/recommend", json=body
    ).json()["data"]["recommendation_id"]

    xlsx = api_client.get(
        f"/api/v1/supplier-recommendations/{rec_id}/export", params={"format": "xlsx"}
    )
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert "Ranked Suppliers" in workbook.sheetnames
    assert "Score Breakdown" in workbook.sheetnames

    csv_response = api_client.get(
        f"/api/v1/supplier-recommendations/{rec_id}/export", params={"format": "csv"}
    )
    assert csv_response.status_code == 200
    assert b"Supplier ID" in csv_response.content

    json_response = api_client.get(
        f"/api/v1/supplier-recommendations/{rec_id}/export", params={"format": "json"}
    )
    assert json_response.status_code == 200
    payload = json.loads(json_response.content)
    assert payload["report_type"] == "sap_supplier_recommendation"
    assert "disclaimer" in payload
