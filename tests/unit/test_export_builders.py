"""Unit tests for the two export builders added to close the Phase 5 gap.

These test the builders directly, from hand-built payloads, so a failure points
at the report rather than at the module that fed it. The API-level and
end-to-end coverage lives in ``tests/api`` and ``tests/e2e``.

What is asserted here is *content*, not formatting. A report that opens is not
the same as a report that says the right thing, and the two things most likely
to go quietly wrong are the ones a spreadsheet cannot show you: a score whose
arithmetic no longer adds up, and a caveat that got dropped on the way out.
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.exports.interview_report_builder import (
    build_interview_csv_report,
    build_interview_json_report,
    build_interview_pdf_report,
    build_interview_xlsx_report,
    interview_disclaimer,
)
from app.services.exports.styling import cell_value, humanise, scalar
from app.services.exports.supplier_risk_report_builder import (
    build_supplier_risk_csv_report,
    build_supplier_risk_json_report,
    build_supplier_risk_xlsx_report,
    supplier_risk_disclaimer,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def risk_payload() -> dict:
    """One assessment with two suppliers: one fully scored, one with holes."""
    return {
        "assessment": {
            "assessment_id": "a1b2c3d4e5f6",
            "dataset_id": "d1d2d3d4",
            "status": "completed",
            "source_filename": "risk_profiles.csv",
            "as_of_date": "2026-07-01",
            "base_currency": "EUR",
            "config_version": "1.0.0",
            "engine_version": "1.0.0",
            "created_at": "2026-08-01T09:00:00Z",
            "supplier_id_filter": None,
        },
        "summary": {"supplier_count": 2, "high_risk_count": 1, "average_score": 55.0},
        "weights": {"delivery": 15.0, "quality": 15.0, "financial": 10.0},
        "suppliers": [
            {
                "rank": 1,
                "supplier_id": "0000390001",
                "supplier_name": "Meridian Components AG",
                "overall_score": 78.5,
                "overall_band": "high",
                "country": "DE",
                "total_spend_base": 1200000.0,
                "data_completeness_pct": 100.0,
                "limited_data": False,
                "trend_direction": "worsening",
                "output_origin": "rule_based",
            },
            {
                "rank": 2,
                "supplier_id": "0000390002",
                "supplier_name": "Calder & Sons Ltd",
                "overall_score": 31.5,
                "overall_band": "low",
                "country": "GB",
                "total_spend_base": 90000.0,
                "data_completeness_pct": 40.0,
                "limited_data": True,
                "trend_direction": "stable",
                "output_origin": "rule_based",
            },
        ],
        "profiles": [
            {
                "supplier_id": "0000390001",
                "supplier_name": "Meridian Components AG",
                "overall_score": 78.5,
                "overall_band": "high",
                "data_completeness_pct": 100.0,
                "limited_data": False,
                "scored_categories": ["delivery", "quality", "financial"],
                "unscored_categories": [],
                "categories": [
                    {
                        "category": "delivery",
                        "label": "Delivery",
                        "score": 80.0,
                        "band": "high",
                        "weight": 15.0,
                        "normalized_weight": 0.375,
                        "contribution": 30.0,
                        "data_available": True,
                        "missing_metrics": [],
                    },
                    {
                        "category": "quality",
                        "label": "Quality",
                        "score": 90.0,
                        "band": "high",
                        "weight": 15.0,
                        "normalized_weight": 0.375,
                        "contribution": 33.75,
                        "data_available": True,
                        "missing_metrics": [],
                    },
                    {
                        "category": "financial",
                        "label": "Financial",
                        "score": 59.0,
                        "band": "medium",
                        "weight": 10.0,
                        "normalized_weight": 0.25,
                        "contribution": 14.75,
                        "data_available": True,
                        "missing_metrics": [],
                    },
                ],
                "actions": [
                    {
                        "category": "delivery",
                        "category_label": "Delivery",
                        "priority": "high",
                        "action": "Book a delivery performance review.",
                        "trigger": "9 late deliveries in the last 90 days",
                        "output_origin": "rule_based",
                    }
                ],
                "delivery_issues": [{"severity": "high", "description": "Late by 14 days"}],
                "quality_issues": [],
                "invoice_issues": [],
                "compliance_issues": [],
                "trend": {"direction": "worsening", "basis": "events in the last 90 days"},
                "event_count": 12,
                "purchase_order_count": 40,
                "invoice_count": 38,
            },
            {
                "supplier_id": "0000390002",
                "supplier_name": "Calder & Sons Ltd",
                "overall_score": 31.5,
                "overall_band": "low",
                "data_completeness_pct": 40.0,
                "limited_data": True,
                "scored_categories": ["delivery"],
                "unscored_categories": ["quality", "financial"],
                "categories": [
                    {
                        "category": "delivery",
                        "label": "Delivery",
                        "score": 31.5,
                        "band": "low",
                        "weight": 15.0,
                        "normalized_weight": 1.0,
                        "contribution": 31.5,
                        "data_available": True,
                        "missing_metrics": [],
                    },
                    {
                        "category": "quality",
                        "label": "Quality",
                        "score": None,
                        "band": None,
                        "weight": 15.0,
                        "normalized_weight": 0.0,
                        "contribution": None,
                        "data_available": False,
                        "missing_metrics": ["defect_rate", "quality_score"],
                    },
                ],
                "actions": [],
                "delivery_issues": [],
                "quality_issues": [],
                "invoice_issues": [],
                "compliance_issues": [],
                "trend": {"direction": "stable", "basis": "no events"},
                "event_count": 0,
                "purchase_order_count": 3,
                "invoice_count": 2,
            },
        ],
        "rule_errors": [],
        "ai_narrative": {"available": True, "origin": "mock_ai", "summary": "Two suppliers."},
        "methodology": {"scale": "0-100, higher is riskier", "weight_total": 100},
    }


@pytest.fixture
def interview_payload() -> dict:
    """A two-answer session: one scored, one left unanswered."""
    return {
        "session": {
            "session_id": "s1s2s3s4s5s6",
            "name": "Demo practice session",
            "candidate_name": "Demo candidate",
            "status": "completed",
            "tracks": ["sap_mm"],
            "mode": "practice",
            "difficulties": ["foundational"],
            "topics": ["Purchasing documents"],
            "requested_question_count": 2,
            "seed": 20260801,
            "question_bank_version": "1.0.0",
            "config_version": "1.0.0",
            "engine_version": "1.0.0",
            "started_at": "2026-08-01T09:00:00Z",
            "completed_at": "2026-08-01T09:20:00Z",
        },
        "summary": {
            "question_count": 2,
            "answered_count": 1,
            "pending_count": 1,
            "average_score": 62.0,
            "band": "developing",
            "passed_count": 1,
            "stale_score_count": 0,
            "over_time_count": 0,
            "total_seconds_spent": 90,
            "average_by_dimension": {"technical_accuracy": 70.0},
            "average_by_topic": {"Purchasing documents": 62.0},
            "average_by_difficulty": {"foundational": 62.0},
            "recommended_study_topics": ["Release strategies"],
            "most_missed_concepts": ["Release code"],
            "strong_topics": [],
            "weak_topics": ["Purchasing documents"],
        },
        "answers": [
            {
                "answer_id": "ans-1",
                "position": 1,
                "status": "answered",
                "answer_text": "A requisition is internal; a purchase order goes to the supplier.",
                "seconds_spent": 90,
                "within_time_limit": True,
                "scoring_is_stale": False,
                "question": {
                    "question_id": "IQ-MM-001",
                    "track": "sap_mm",
                    "topic": "Purchasing documents",
                    "difficulty": "foundational",
                    "question": "What is the difference between a requisition and a purchase order?",
                },
                "score": {
                    "overall_score": 62.0,
                    "overall_band": "developing",
                    "passed": True,
                    "pass_score": 60.0,
                    "output_origin": "rule_based",
                    "dimensions": [
                        {
                            "dimension": "technical_accuracy",
                            "score": 70.0,
                            "weight": 0.6,
                            "applicable": True,
                            "band": "developing",
                            "explanation": "2 of 3 expected concept(s) covered.",
                            "concepts_expected": 3,
                            "concepts_matched": 2,
                        },
                        {
                            "dimension": "architecture",
                            "score": None,
                            "weight": 0.0,
                            "applicable": False,
                            "band": None,
                            "explanation": "This question tests no architecture concept.",
                            "concepts_expected": 0,
                            "concepts_matched": 0,
                        },
                    ],
                    "concept_matches": [
                        {
                            "concept_id": "pr_internal",
                            "label": "A requisition is an internal request",
                            "dimension": "technical",
                            "weight": 1.5,
                            "required": True,
                            "matched": True,
                            "matched_keyword": "internal",
                            "excerpt": "A requisition is internal",
                            "negated": False,
                            "study_hint": "Say the requisition never reaches the supplier.",
                        },
                        {
                            "concept_id": "release_code",
                            "label": "A release strategy uses release codes",
                            "dimension": "technical",
                            "weight": 1.0,
                            "required": False,
                            "matched": False,
                            "matched_keyword": None,
                            "excerpt": None,
                            "negated": False,
                            "study_hint": "Mention the release code.",
                        },
                    ],
                },
                "feedback": {
                    "strengths": ["Named the internal/external split clearly."],
                    "missing_concepts": ["Release code"],
                    "incorrect_statements": ["A requisition is sent to the supplier."],
                    "improved_sample_answer": "A purchase requisition is an internal request...",
                    "topics_to_study": ["Release strategies"],
                    "coaching_note": "Good start. Name the release code next time.",
                    "source": "ai_generated",
                    "output_origin": "mock_ai",
                    "ai_provider": "mock",
                },
            },
            {
                "answer_id": "ans-2",
                "position": 2,
                "status": "pending",
                "answer_text": None,
                "seconds_spent": None,
                "within_time_limit": True,
                "scoring_is_stale": False,
                "question": {
                    "question_id": "IQ-MM-002",
                    "track": "sap_mm",
                    "topic": "Purchasing documents",
                    "difficulty": "foundational",
                    "question": "What does a release strategy do?",
                },
                "score": None,
                "feedback": None,
            },
        ],
        "notes": [],
        "uncovered_tracks": [],
        "methodology": {"pass_score": 60.0, "disclaimer": "Practice feedback."},
    }


def _sheet_text(content: bytes, name: str) -> str:
    """Every cell of one sheet, flattened, for a substring check."""
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook[name]
    return "\n".join(
        " | ".join("" if cell.value is None else str(cell.value) for cell in row)
        for row in sheet.iter_rows()
    )


# ---------------------------------------------------------------------------
# Shared styling helpers
# ---------------------------------------------------------------------------
class TestStyling:
    def test_a_boolean_prints_as_a_word_not_as_true(self):
        """Every existing builder does this; a shared helper must not differ."""
        assert scalar(True) == "Yes"
        assert scalar(False) == "No"

    def test_a_structure_is_serialised_rather_than_stringified(self):
        assert json.loads(scalar({"a": 1})) == {"a": 1}

    def test_a_long_structure_is_truncated_to_fit_a_cell(self):
        assert len(scalar({str(index): "x" * 100 for index in range(200)})) <= 2000

    def test_a_list_of_short_values_reads_as_a_pipe_list(self):
        assert cell_value({"k": ["a", "b"]}, "k") == "a | b"

    def test_a_missing_key_is_empty_rather_than_an_error(self):
        assert cell_value({}, "nope") is None

    @pytest.mark.parametrize(
        ("key", "expected"),
        [
            ("overall_score", "Overall score"),
            ("total_spend_base", "Total spend (base)"),
            ("spend_pct", "Spend %"),
        ],
    )
    def test_a_field_name_becomes_a_readable_label(self, key, expected):
        assert humanise(key) == expected


# ---------------------------------------------------------------------------
# Module 5
# ---------------------------------------------------------------------------
class TestSupplierRiskReport:
    def test_the_json_report_carries_everything_the_api_had(self, risk_payload):
        document = json.loads(build_supplier_risk_json_report(risk_payload))
        assert document["report_type"] == "sap_supplier_risk_assessment"
        assert document["disclaimer"]
        for key in (
            "assessment", "summary", "weights", "suppliers", "profiles",
            "rule_errors", "ai_narrative", "methodology",
        ):
            assert key in document, f"the JSON report dropped {key}"
        assert document["assessment"]["assessment_id"] == "a1b2c3d4e5f6"
        assert len(document["profiles"]) == 2

    def test_the_csv_has_one_row_per_supplier_and_a_header(self, risk_payload):
        text = build_supplier_risk_csv_report(risk_payload["suppliers"]).decode("utf-8-sig")
        rows = [line for line in text.splitlines() if line.strip()]
        assert len(rows) == 3
        assert "Supplier ID" in rows[0] and "Overall Risk Score" in rows[0]
        assert "0000390001" in rows[1]

    def test_the_csv_opens_in_excel(self, risk_payload):
        """The BOM is what stops Excel mangling a UTF-8 supplier name."""
        assert build_supplier_risk_csv_report(risk_payload["suppliers"]).startswith(
            b"\xef\xbb\xbf"
        )

    def test_the_workbook_has_the_six_documented_sheets(self, risk_payload):
        workbook = load_workbook(io.BytesIO(build_supplier_risk_xlsx_report(risk_payload)))
        assert workbook.sheetnames == [
            "Summary",
            "Portfolio",
            "Category Scores",
            "Evidence & Missing Data",
            "Recommended Actions",
            "Methodology",
        ]

    def test_the_contributions_reconstruct_the_overall_score(self, risk_payload):
        """The point of printing the arithmetic is that it adds up.

        30.0 + 33.75 + 14.75 = 78.5. A reader who cannot re-derive a risk score
        will argue with it rather than act on it.
        """
        text = _sheet_text(build_supplier_risk_xlsx_report(risk_payload), "Category Scores")
        assert "30" in text and "33.75" in text and "14.75" in text
        contributions = [
            category["contribution"]
            for category in risk_payload["profiles"][0]["categories"]
        ]
        assert sum(contributions) == pytest.approx(
            risk_payload["profiles"][0]["overall_score"]
        )

    def test_a_category_with_no_data_is_shown_as_unscored_not_as_zero(self, risk_payload):
        """A 31.5 built from one category must not look like a 31.5 built from ten."""
        content = build_supplier_risk_xlsx_report(risk_payload)
        categories = _sheet_text(content, "Category Scores")
        assert "No" in categories, "the data-available column is missing"
        assert "defect_rate" in categories, "the missing metrics are not listed"

        evidence = _sheet_text(content, "Evidence & Missing Data")
        assert "quality" in evidence and "financial" in evidence
        assert "40" in evidence, "the data completeness figure is missing"

    def test_the_recommended_actions_name_what_triggered_them(self, risk_payload):
        text = _sheet_text(build_supplier_risk_xlsx_report(risk_payload), "Recommended Actions")
        assert "Book a delivery performance review." in text
        assert "9 late deliveries in the last 90 days" in text

    def test_the_methodology_sheet_carries_the_weights_and_the_disclaimer(self, risk_payload):
        text = _sheet_text(build_supplier_risk_xlsx_report(risk_payload), "Methodology")
        assert "Delivery" in text and "15" in text
        assert "NO EXTERNAL DATA SOURCE" in text

    def test_every_format_carries_the_external_data_caveat(self, risk_payload):
        """The caveat people forget: a low score is a clean *internal* record."""
        assert "NO EXTERNAL DATA SOURCE WAS CONSULTED" in supplier_risk_disclaimer()
        assert "credit bureau" in supplier_risk_disclaimer()

        document = json.loads(build_supplier_risk_json_report(risk_payload))
        assert "NO EXTERNAL DATA SOURCE" in document["disclaimer"]
        assert "NO EXTERNAL DATA SOURCE" in _sheet_text(
            build_supplier_risk_xlsx_report(risk_payload), "Summary"
        )

    def test_an_empty_assessment_still_produces_a_readable_file(self):
        """No suppliers is a state, not an error - the sheets and caveat remain."""
        empty = {"assessment": {"assessment_id": "x"}, "suppliers": [], "profiles": []}
        workbook = load_workbook(io.BytesIO(build_supplier_risk_xlsx_report(empty)))
        assert len(workbook.sheetnames) == 6
        assert build_supplier_risk_csv_report([]).decode("utf-8-sig").strip()
        assert json.loads(build_supplier_risk_json_report(empty))["disclaimer"]


# ---------------------------------------------------------------------------
# Module 10
# ---------------------------------------------------------------------------
class TestInterviewReport:
    def test_the_json_report_carries_everything_the_api_had(self, interview_payload):
        document = json.loads(build_interview_json_report(interview_payload))
        assert document["report_type"] == "sap_interview_session"
        for key in ("session", "summary", "answers", "methodology", "disclaimer"):
            assert key in document, f"the JSON report dropped {key}"
        assert len(document["answers"]) == 2

    def test_the_csv_has_one_row_per_answer(self, interview_payload):
        text = build_interview_csv_report(interview_payload["answers"]).decode("utf-8-sig")
        rows = [line for line in text.splitlines() if line.strip()]
        assert rows[0].startswith("#,Question ID")
        assert "IQ-MM-001" in text and "IQ-MM-002" in text

    def test_the_workbook_has_the_five_documented_sheets(self, interview_payload):
        workbook = load_workbook(io.BytesIO(build_interview_xlsx_report(interview_payload)))
        assert workbook.sheetnames == [
            "Summary",
            "Answers",
            "Dimension Scores",
            "Concept Coverage",
            "Study Plan",
        ]

    def test_the_answers_sheet_carries_the_question_and_what_was_written(
        self, interview_payload
    ):
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Answers")
        assert "What is the difference between a requisition and a purchase order?" in text
        assert "A requisition is internal; a purchase order goes to the supplier." in text
        assert "Named the internal/external split clearly." in text
        assert "A requisition is sent to the supplier." in text, "corrections were dropped"
        assert "A purchase requisition is an internal request..." in text

    def test_a_dimension_that_does_not_apply_is_not_a_zero(self, interview_payload):
        """`architecture` is `applicable: false`, not `score: 0`."""
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Dimension Scores")
        assert "architecture" in text
        assert "This question tests no architecture concept." in text
        assert "No" in text, "the applicable column is missing"

    def test_a_credited_concept_says_what_credited_it(self, interview_payload):
        """Keyword matching has a ceiling, so the decision has to be inspectable."""
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Concept Coverage")
        assert "A requisition is an internal request" in text
        assert "internal" in text, "the matched keyword is missing"
        assert "Mention the release code." in text, "the study hint is missing"

    def test_the_study_sheet_carries_the_recommended_topics(self, interview_payload):
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Study Plan")
        assert "Release strategies" in text
        assert "Release code" in text

    def test_the_summary_separates_unanswered_from_zero(self, interview_payload):
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Summary")
        assert "Pending count" in text or "pending" in text.lower()

    def test_the_prose_carries_its_origin_and_the_score_carries_its_own(
        self, interview_payload
    ):
        """A candidate must be able to tell a mark from a paragraph."""
        text = _sheet_text(build_interview_xlsx_report(interview_payload), "Answers")
        assert "mock_ai" in text, "the feedback origin is missing"
        assert "rule_based" in text, "the score origin is missing"

    def test_the_pdf_is_a_readable_transcript(self, interview_payload):
        content = build_interview_pdf_report(interview_payload)
        assert content.startswith(b"%PDF-")

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "SAP INTERVIEW COACH" in text
        assert "IQ-MM-001" in text
        assert "Candidate answer" in text
        assert "credited by" in text, "the matched keyword is missing from the transcript"
        assert "PRACTICE FEEDBACK" in text, "the disclaimer is missing from the transcript"

    def test_the_pdf_gives_each_answer_its_own_page(self, interview_payload):
        reader = PdfReader(io.BytesIO(build_interview_pdf_report(interview_payload)))
        # A header page plus one page per answer, at minimum.
        assert len(reader.pages) >= 1 + len(interview_payload["answers"])

    def test_an_unanswered_question_is_reported_not_scored(self, interview_payload):
        reader = PdfReader(io.BytesIO(build_interview_pdf_report(interview_payload)))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "IQ-MM-002" in text
        assert "not counted as zero" in text

    def test_every_format_says_this_is_practice_not_an_assessment(self, interview_payload):
        assert "NOT AN ASSESSMENT" in interview_disclaimer()
        assert "certification" in interview_disclaimer()

        document = json.loads(build_interview_json_report(interview_payload))
        assert "NOT AN ASSESSMENT" in document["disclaimer"]
        assert "NOT AN ASSESSMENT" in _sheet_text(
            build_interview_xlsx_report(interview_payload), "Summary"
        )

    def test_a_session_with_no_answers_still_produces_a_readable_file(self):
        empty = {"session": {"session_id": "x", "name": "Empty"}, "summary": {}, "answers": []}
        workbook = load_workbook(io.BytesIO(build_interview_xlsx_report(empty)))
        assert len(workbook.sheetnames) == 5
        assert build_interview_pdf_report(empty).startswith(b"%PDF-")
        assert json.loads(build_interview_json_report(empty))["disclaimer"]
