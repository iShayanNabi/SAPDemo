"""One end-to-end workflow per module, driven over HTTP against the demo data.

``tests/api`` proves a route works. These prove a *journey* works: the
identifier one call issues is accepted by the next, the list reflects what was
just created, the detail response and the downloaded report agree, and every
step returns the shared envelope.

They deliberately use the bundled sample datasets rather than hand-built
fixtures. That is the same path a first-time user takes from the README, so a
sample file that stops matching its module fails here rather than in front of
somebody following the walkthrough.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.api.test_blueprint_api import PROJECT
from tests.api.test_supplier_reco_api import REQUIREMENT
from tests.api.test_test_case_api import PROCESS_CONTEXT
from tests.e2e.conftest import (
    CSV_MIME,
    PDF_MIME,
    VALID_ORIGINS,
    assert_identifier,
    assert_pagination,
    assert_timestamp,
    download,
    failure,
    ok,
    upload,
)

pytestmark = pytest.mark.slow

V1 = "/api/v1"


def _sample(name: str) -> Path:
    """Return a bundled demo file, skipping the test when it is not generated."""
    from app.core.config import PROJECT_ROOT

    path = PROJECT_ROOT / "data" / "sample" / name
    if not path.is_file():
        pytest.skip(f"{name} is missing - run the matching scripts/generate_*.py")
    return path


# ===========================================================================
# Module 1 - Purchase Order Risk Checker
# ===========================================================================
class TestPurchaseOrderRiskWorkflow:
    """Upload a PO export, analyse it, read the findings, download the report."""

    def test_the_whole_journey(self, client):
        # 1. The user checks the lab is up before uploading anything.
        health = ok(client.get(f"{V1}/health"))
        assert health["status"] == "ok"
        assert health["database_connected"] is True

        # 2. They look at the bundled dataset and download it.
        info = ok(client.get(f"{V1}/po-risk/sample/info"))
        assert info["available"] is True
        assert info["row_count"] > 0
        sample_bytes = download(
            client.get(f"{V1}/po-risk/sample", params={"format": "csv"}), expect_extension="csv"
        )
        assert sample_bytes[:1] != b"PK", "the CSV sample must not be a workbook"

        # 3. Upload. The mapping is suggested, not demanded.
        uploaded = upload(
            client, sample_bytes, f"{V1}/po-risk/upload",
            filename="sample_purchase_orders.csv", mime=CSV_MIME,
        )
        upload_id = assert_identifier(uploaded["upload_id"], label="upload_id")
        assert uploaded["is_analyzable"] is True
        assert uploaded["missing_required_fields"] == []
        assert uploaded["row_count"] == info["row_count"]
        assert uploaded["suggested_mapping"], "no column mapping was suggested"

        # 4. Analyse.
        analysis = ok(
            client.post(
                f"{V1}/po-risk/analyze",
                json={"upload_id": upload_id, "generate_ai_summary": True},
            )
        )
        analysis_id = assert_identifier(analysis["analysis_id"], label="analysis_id")
        assert analysis["status"] == "completed"
        assert analysis["upload_id"] == upload_id
        assert analysis["findings_count"] > 0
        assert analysis["rule_errors"] == [], f"a rule failed: {analysis['rule_errors']}"
        assert_timestamp(analysis["created_at"], label="analysis.created_at")
        assert_timestamp(analysis["completed_at"], label="analysis.completed_at")

        # The AI narrative is optional enrichment in its own field and is
        # labelled as mock, never merged into a computed figure.
        assert analysis["ai_output_origin"] == "mock_ai"
        assert isinstance(analysis["risk_score"], (int, float))

        # The severity counts must add up to the findings count - two fields
        # describing the same thing, asserted together.
        counted = sum(
            analysis[f"{level}_count"] for level in ("low", "medium", "high", "critical")
        )
        assert counted == analysis["findings_count"]

        # 5. The analysis appears in the list.
        listed = ok(client.get(f"{V1}/po-risk/analyses", params={"limit": 5}))
        assert_pagination(listed, items_key="analyses")
        assert analysis_id in {row["analysis_id"] for row in listed["analyses"]}

        # 6. Re-reading it gives the same figures.
        fetched = ok(client.get(f"{V1}/po-risk/analyses/{analysis_id}"))
        assert fetched["findings_count"] == analysis["findings_count"]
        assert fetched["risk_score"] == analysis["risk_score"]

        # 7. Findings: paginated, filterable, and every one labelled.
        page = ok(
            client.get(
                f"{V1}/po-risk/analyses/{analysis_id}/findings", params={"limit": 25, "offset": 0}
            )
        )
        assert_pagination(page, items_key="findings")
        assert page["total"] == analysis["findings_count"]
        for finding in page["findings"]:
            assert finding["output_origin"] in VALID_ORIGINS
            assert finding["severity"] in {"low", "medium", "high", "critical"}
            assert 0.0 <= finding["confidence_score"] <= 1.0
            assert finding["rule_id"].startswith("PO-R")
            assert finding["explanation"]

        critical = ok(
            client.get(
                f"{V1}/po-risk/analyses/{analysis_id}/findings", params={"severity": "critical"}
            )
        )
        assert all(item["severity"] == "critical" for item in critical["findings"])
        assert critical["total"] == analysis["critical_count"]

        # A second page must not repeat the first.
        if page["total"] > 25:
            second = ok(
                client.get(
                    f"{V1}/po-risk/analyses/{analysis_id}/findings",
                    params={"limit": 25, "offset": 25},
                )
            )
            first_ids = {item["finding_id"] for item in page["findings"]}
            assert not first_ids & {item["finding_id"] for item in second["findings"]}

        # 8. Download the report in all three formats.
        workbook_bytes = download(
            client.get(f"{V1}/po-risk/analyses/{analysis_id}/export", params={"format": "xlsx"}),
            expect_extension="xlsx",
        )
        workbook = load_workbook(io.BytesIO(workbook_bytes))
        assert workbook.sheetnames, "the workbook has no sheets"

        download(
            client.get(f"{V1}/po-risk/analyses/{analysis_id}/export", params={"format": "csv"}),
            expect_extension="csv",
        )
        report = json.loads(
            download(
                client.get(
                    f"{V1}/po-risk/analyses/{analysis_id}/export", params={"format": "json"}
                ),
                expect_extension="json",
            )
        )
        assert report["analysis"]["analysis_id"] == analysis_id
        assert len(report["findings"]) == analysis["findings_count"]

        # 9. An unknown analysis is a clean 404, not a stack trace.
        error = failure(client.get(f"{V1}/po-risk/analyses/does-not-exist"), expected_status=404)
        assert error["code"] == "not_found"
        assert "Traceback" not in error["message"]


# ===========================================================================
# Module 2 - Spend Analytics Dashboard
# ===========================================================================
class TestSpendAnalyticsWorkflow:
    """Upload spend, build the dashboard, drill down, download the report."""

    def test_the_whole_journey(self, client):
        uploaded = upload(
            client, _sample("sample_spend_transactions.csv"), f"{V1}/spend/upload",
            filename="sample_spend_transactions.csv", mime=CSV_MIME,
        )
        assert uploaded["is_analyzable"] is True

        analysis = ok(
            client.post(
                f"{V1}/spend/analyze",
                json={"upload_id": uploaded["upload_id"], "generate_ai_summary": True},
            )
        )
        analysis_id = assert_identifier(analysis["analysis_id"], label="analysis_id")
        assert analysis["status"] == "completed"
        assert analysis["record_count"] > 0
        assert analysis["savings_errors"] == []
        metrics = analysis["metrics"]
        assert metrics["total_spend"] > 0
        assert metrics["output_origin"] == "rule_based", "no spend figure is written by AI"
        # Contracted plus non-contracted is the whole of the spend: two fields
        # describing the same money, asserted together.
        assert metrics["contracted_spend"] + metrics["non_contracted_spend"] == pytest.approx(
            metrics["total_spend"], rel=1e-6
        )

        # The dashboard the UI draws.
        assert analysis["analytics"], "no breakdowns were produced"

        listed = ok(client.get(f"{V1}/spend/analyses", params={"limit": 5}))
        assert_pagination(listed, items_key="analyses")
        assert analysis_id in {row["analysis_id"] for row in listed["analyses"]}

        # Drill-down: click a supplier in the chart, get exactly its lines.
        top_supplier = analysis["supplier_spend"][0]
        drill = ok(
            client.get(
                f"{V1}/spend/analyses/{analysis_id}/transactions",
                params={"supplier_id": top_supplier["supplier_id"], "limit": 50},
            )
        )
        assert_pagination(drill, items_key="transactions")
        assert drill["total"] > 0
        assert all(
            row["supplier_id"] == top_supplier["supplier_id"] for row in drill["transactions"]
        )
        # The drill-down total must equal the figure the chart element showed.
        # This is the pair that reported 329,444,859 EUR for 7 transactions.
        assert drill["total_spend_base"] == pytest.approx(top_supplier["spend_base"], rel=1e-6)

        # Savings opportunities are estimates and say so.
        opportunities = ok(
            client.get(
                f"{V1}/spend/analyses/{analysis_id}/opportunities", params={"limit": 20}
            )
        )
        assert_pagination(opportunities, items_key="opportunities")
        assert opportunities["disclaimer"]
        for opportunity in opportunities["opportunities"]:
            assert opportunity["is_estimate"] is True
            assert opportunity["output_origin"] in VALID_ORIGINS
            assert 0.0 <= opportunity["confidence"] <= 1.0

        download(
            client.get(f"{V1}/spend/analyses/{analysis_id}/export", params={"format": "xlsx"}),
            expect_extension="xlsx",
        )
        report = json.loads(
            download(
                client.get(f"{V1}/spend/analyses/{analysis_id}/export", params={"format": "json"}),
                expect_extension="json",
            )
        )
        assert report["analysis"]["analysis_id"] == analysis_id


# ===========================================================================
# Module 3 - Supplier Recommendation Engine
# ===========================================================================
class TestSupplierRecommendationWorkflow:
    """Load a catalogue, ask for a requirement, read the ranking, export it."""

    def test_the_whole_journey(self, client):
        catalogue = upload(
            client, _sample("sample_suppliers.csv"), f"{V1}/suppliers/upload",
            filename="sample_suppliers.csv", mime=CSV_MIME,
        )
        catalog_id = assert_identifier(catalogue["catalog_id"], label="catalog_id")
        assert catalogue["supplier_count"] > 0

        suppliers = ok(client.get(f"{V1}/suppliers", params={"limit": 5}))
        assert_pagination(suppliers, items_key="suppliers")
        first_id = suppliers["suppliers"][0]["supplier_id"]
        profile = ok(client.get(f"{V1}/suppliers/{first_id}"))
        assert profile["supplier_id"] == first_id

        recommendation = ok(
            client.post(
                f"{V1}/supplier-recommendations/recommend",
                json={
                    "catalog_id": catalog_id,
                    "requirement": REQUIREMENT,
                    "generate_ai_summary": True,
                },
            )
        )
        recommendation_id = assert_identifier(
            recommendation["recommendation_id"], label="recommendation_id"
        )
        assert recommendation["status"] == "completed"
        assert recommendation["disclaimer"], "a ranking must carry its standing caveat"
        results = recommendation["results"]
        assert results, "no supplier was ranked"

        # Eligible and ineligible must account for the whole catalogue: two
        # counts describing one set, asserted together.
        assert (
            recommendation["eligible_count"] + recommendation["ineligible_count"]
            == recommendation["total_supplier_count"]
        )

        eligible = [entry for entry in results if entry["is_eligible"]]
        assert len(eligible) == recommendation["eligible_count"]

        # Ranking must be a real ordering over the eligible suppliers, and each
        # score has to be explainable from its own components.
        scores = [entry["overall_score"] for entry in eligible]
        assert scores == sorted(scores, reverse=True)
        assert [entry["rank"] for entry in eligible] == list(range(1, len(eligible) + 1))
        for entry in eligible:
            assert entry["output_origin"] in VALID_ORIGINS
            assert entry["evidence"], "a score with no evidence is not explainable"
            assert entry["explanation"]
        for entry in results:
            if not entry["is_eligible"]:
                assert entry["ineligibility_reasons"], (
                    "a supplier excluded for no stated reason cannot be argued with"
                )

        listed = ok(client.get(f"{V1}/supplier-recommendations", params={"limit": 5}))
        assert_pagination(listed, items_key="recommendations")
        assert recommendation_id in {row["recommendation_id"] for row in listed["recommendations"]}

        fetched = ok(client.get(f"{V1}/supplier-recommendations/{recommendation_id}"))
        assert [entry["supplier_id"] for entry in fetched["results"]] == [
            entry["supplier_id"] for entry in results
        ]

        download(
            client.get(
                f"{V1}/supplier-recommendations/{recommendation_id}/export",
                params={"format": "xlsx"},
            ),
            expect_extension="xlsx",
        )


# ===========================================================================
# Module 4 - Invoice Validator
# ===========================================================================
class TestInvoiceValidatorWorkflow:
    """Three-way match: invoices against purchase orders and goods receipts."""

    def test_the_whole_journey(self, client):
        def load(name: str, dataset: str) -> dict:
            path = _sample(name)
            return ok(
                client.post(
                    f"{V1}/invoices/upload",
                    data={"dataset": dataset},
                    files={"file": (path.name, path.read_bytes(), CSV_MIME)},
                )
            )

        invoices = load("sample_invoices.csv", "invoices")
        purchase_orders = load("sample_invoice_purchase_orders.csv", "purchase_orders")
        receipts = load("sample_goods_receipts.csv", "goods_receipts")
        # Each upload knows which of the three datasets it is.
        assert {invoices["dataset"], purchase_orders["dataset"], receipts["dataset"]} == {
            "invoices", "purchase_orders", "goods_receipts",
        }

        validation = ok(
            client.post(
                f"{V1}/invoices/validate",
                json={
                    "invoice_upload_id": invoices["upload_id"],
                    "po_upload_id": purchase_orders["upload_id"],
                    "gr_upload_id": receipts["upload_id"],
                    "generate_ai_summary": True,
                },
            )
        )
        validation_id = assert_identifier(validation["validation_id"], label="validation_id")
        assert validation["status"] == "completed"
        assert validation["exceptions_count"] > 0
        assert validation["rule_errors"] == []

        listed = ok(client.get(f"{V1}/invoices/validations", params={"limit": 5}))
        assert_pagination(listed, items_key="validations")
        assert validation_id in {row["validation_id"] for row in listed["validations"]}

        exceptions = ok(
            client.get(
                f"{V1}/invoices/validations/{validation_id}/exceptions", params={"limit": 25}
            )
        )
        assert_pagination(exceptions, items_key="exceptions")
        assert exceptions["total"] == validation["exceptions_count"]
        for item in exceptions["exceptions"]:
            assert item["severity"] in {"low", "medium", "high", "critical"}
            assert item["output_origin"] in VALID_ORIGINS
            assert item["rule_id"].startswith("IV-R")

        # Filtering by rule must not change what a rule means.
        rule_id = exceptions["exceptions"][0]["rule_id"]
        filtered = ok(
            client.get(
                f"{V1}/invoices/validations/{validation_id}/exceptions",
                params={"rule_id": rule_id},
            )
        )
        assert all(item["rule_id"] == rule_id for item in filtered["exceptions"])
        assert 0 < filtered["total"] <= exceptions["total"]

        download(
            client.get(
                f"{V1}/invoices/validations/{validation_id}/export", params={"format": "xlsx"}
            ),
            expect_extension="xlsx",
        )


# ===========================================================================
# Module 5 - Supplier Risk Copilot
# ===========================================================================
class TestSupplierRiskWorkflow:
    """Score a portfolio, read one profile, then ask the copilot about it."""

    def test_the_whole_journey(self, client):
        profiles = upload(
            client, _sample("sample_supplier_risk_profiles.csv"), f"{V1}/supplier-risk/upload",
            filename="sample_supplier_risk_profiles.csv", mime=CSV_MIME,
        )
        dataset_id = assert_identifier(profiles["dataset_id"], label="dataset_id")
        assert profiles["is_analyzable"] is True
        assert profiles["supplier_count"] > 0

        # Risk events attach to the same dataset.
        events_path = _sample("sample_supplier_risk_events.csv")
        events = ok(
            client.post(
                f"{V1}/supplier-risk/upload",
                data={"dataset": "events", "dataset_id": dataset_id},
                files={"file": (events_path.name, events_path.read_bytes(), CSV_MIME)},
            )
        )
        assert events["dataset_id"] == dataset_id
        assert events["event_count"] > 0

        assessment = ok(
            client.post(
                f"{V1}/supplier-risk/calculate",
                json={"dataset_id": dataset_id, "generate_ai_summary": True},
            )
        )
        assessment_id = assert_identifier(assessment["assessment_id"], label="assessment_id")
        assert assessment["status"] == "completed"
        assert assessment["rule_errors"] == []
        assert assessment["disclaimer"]
        assert len(assessment["suppliers"]) == profiles["supplier_count"]

        listed = ok(client.get(f"{V1}/supplier-risk/assessments", params={"limit": 5}))
        assert_pagination(listed, items_key="assessments")
        assert assessment_id in {row["assessment_id"] for row in listed["assessments"]}

        portfolio = ok(client.get(f"{V1}/supplier-risk/suppliers", params={"limit": 10}))
        assert_pagination(portfolio, items_key="suppliers")
        riskiest = portfolio["suppliers"][0]
        assert riskiest["overall_band"]
        assert riskiest["rank"] == 1

        profile = ok(client.get(f"{V1}/supplier-risk/suppliers/{riskiest['supplier_id']}"))
        assert profile["supplier_id"] == riskiest["supplier_id"]
        assert profile["categories"], "a risk score with no breakdown is not transparent"
        assert profile["overall_score"] == pytest.approx(riskiest["overall_score"])
        # Missing data is reported, never invented: every category is either
        # scored or explicitly listed as unscored.
        assert set(profile["scored_categories"]) | set(profile["unscored_categories"])

        # The copilot answers from the loaded records, with citations.
        answer = ok(
            client.post(
                f"{V1}/supplier-risk/chat",
                json={
                    "question": f"Why is supplier {riskiest['supplier_id']} rated the way it is?",
                    "supplier_id": riskiest["supplier_id"],
                },
            )
        )
        assert answer["answer"]
        assert answer["output_origin"] in VALID_ORIGINS
        assert answer["citations"], "a copilot answer with no citation is not checkable"

        # A question the loaded records cannot answer is refused, not invented.
        refused = ok(
            client.post(
                f"{V1}/supplier-risk/chat",
                json={"question": "What is the share price of this supplier's parent company?"},
            )
        )
        assert refused["answer"]
        assert refused["data_available"] is False
        assert refused["unavailable_reason"], (
            "a copilot that cannot answer must say why, not produce a confident guess"
        )

        # Download the assessment in all three formats.
        workbook_bytes = download(
            client.get(
                f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                params={"format": "xlsx"},
            ),
            expect_extension="xlsx",
        )
        workbook = load_workbook(io.BytesIO(workbook_bytes))
        assert "Methodology" in workbook.sheetnames, (
            "a risk report that does not say how it scored is not defensible"
        )

        download(
            client.get(
                f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                params={"format": "csv"},
            ),
            expect_extension="csv",
        )
        report = json.loads(
            download(
                client.get(
                    f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                    params={"format": "json"},
                ),
                expect_extension="json",
            )
        )
        assert report["assessment"]["assessment_id"] == assessment_id
        assert len(report["suppliers"]) == len(assessment["suppliers"])
        assert "NO EXTERNAL DATA SOURCE" in report["disclaimer"], (
            "the report must say no credit bureau or ESG agency was consulted"
        )
        # The report and the detail response must agree about the same supplier.
        exported = next(
            item for item in report["profiles"] if item["supplier_id"] == profile["supplier_id"]
        )
        assert exported["overall_score"] == profile["overall_score"]
        assert exported["unscored_categories"] == profile["unscored_categories"]

        # One supplier can be exported on its own; the same route, filtered.
        single = json.loads(
            download(
                client.get(
                    f"{V1}/supplier-risk/assessments/{assessment_id}/export",
                    params={"format": "json", "supplier_id": profile["supplier_id"]},
                ),
                expect_extension="json",
            )
        )
        assert [item["supplier_id"] for item in single["profiles"]] == [profile["supplier_id"]]

        # An unknown assessment is a clean 404, not a stack trace.
        error = failure(
            client.get(f"{V1}/supplier-risk/assessments/does-not-exist/export"),
            expected_status=404,
        )
        assert error["code"] == "not_found"
        assert "Traceback" not in error["message"]


# ===========================================================================
# Module 6 - Contract Assistant
# ===========================================================================
class TestContractAssistantWorkflow:
    """Upload a contract document, analyse it, cite it, question it, export it."""

    def test_the_whole_journey(self, client):
        uploaded = upload(
            client, _sample("sample_contract_msa_nordwind.pdf"), f"{V1}/contracts/upload",
            filename="sample_contract_msa_nordwind.pdf", mime=PDF_MIME,
        )
        contract_id = assert_identifier(uploaded["contract_id"], label="contract_id")
        assert uploaded["is_analyzable"] is True
        assert uploaded["page_count"] > 0
        assert uploaded["status"] == "uploaded"

        analysed = ok(
            client.post(
                f"{V1}/contracts/{contract_id}/analyze", json={"generate_ai_summary": True}
            )
        )
        assert analysed["status"] == "analyzed"
        assert analysed["extraction"]["needs_ocr"] is False
        assert analysed["rule_errors"] == []
        assert analysed["clauses"], "no clause was extracted"

        # Every clause that was found carries a source it can be checked
        # against; a clause that is absent says so instead of inventing a page.
        found = [clause for clause in analysed["clauses"] if clause["present"]]
        assert found, "no clause was found in the demo master services agreement"
        for clause in found:
            assert 0.0 <= clause["confidence"] <= 1.0
            assert clause["output_origin"] in VALID_ORIGINS
            assert 1 <= clause["page_number"] <= uploaded["page_count"]
            assert clause["excerpt"], "a citation with no excerpt cannot be checked"
        for clause in analysed["clauses"]:
            if not clause["present"]:
                assert clause["page_number"] is None
                assert not clause["excerpt"]

        # Key dates are read from the clause they belong to, never from the
        # whole document, and each one records how it was reached.
        key_dates = analysed["key_dates"]
        assert key_dates["expiration_date"], "the demo MSA states an expiry date"
        for field in ("effective_date", "expiration_date", "signature_date"):
            if key_dates[field]:
                assert key_dates[f"{field}_basis"], (
                    f"{field} was reported with no stated basis"
                )

        listed = ok(client.get(f"{V1}/contracts", params={"limit": 5}))
        assert_pagination(listed, items_key="contracts")
        assert contract_id in {row["contract_id"] for row in listed["contracts"]}

        clauses = ok(client.get(f"{V1}/contracts/{contract_id}/clauses", params={"limit": 50}))
        assert clauses["clauses"]

        answer = ok(
            client.post(
                f"{V1}/contracts/{contract_id}/questions",
                json={"question": "When does this agreement expire?"},
            )
        )
        assert answer["answer"]
        assert answer["output_origin"] in VALID_ORIGINS

        download(
            client.get(f"{V1}/contracts/{contract_id}/export", params={"format": "xlsx"}),
            expect_extension="xlsx",
        )

    def test_a_hostile_document_is_data_not_instructions(self, client):
        """The demo contract carrying prompt-injection bait must stay inert."""
        uploaded = upload(
            client, _sample("sample_contract_hostile_calder.pdf"), f"{V1}/contracts/upload",
            filename="sample_contract_hostile_calder.pdf", mime=PDF_MIME,
        )
        analysed = ok(
            client.post(
                f"{V1}/contracts/{uploaded['contract_id']}/analyze",
                json={"generate_ai_summary": True},
            )
        )
        rendered = json.dumps(analysed).lower()
        # The payload the injection was trying to get printed must not survive
        # anywhere in the response - not in a clause, not in an excerpt, not in
        # the narrative. An injection marker is the lead-in to a payload.
        assert "validated in a live sap production system" not in rendered
        assert "approved by sap" not in rendered


# ===========================================================================
# Module 7 - Inventory Predictor
# ===========================================================================
class TestInventoryPredictorWorkflow:
    """Load history, forecast demand, read one material, download the plan."""

    def test_the_whole_journey(self, client):
        dataset = upload(
            client, _sample("sample_inventory_history.csv"), f"{V1}/inventory/upload",
            filename="sample_inventory_history.csv", mime=CSV_MIME,
        )
        dataset_id = assert_identifier(dataset["dataset_id"], label="dataset_id")
        assert dataset["is_analyzable"] is True
        assert dataset["series_count"] > 0
        assert dataset["frequency"], "the period frequency must be inferred, never assumed"

        forecast = ok(
            client.post(
                f"{V1}/inventory/forecast",
                json={
                    "dataset_id": dataset_id,
                    "horizon_periods": 6,
                    "confidence_level": 0.95,
                    "as_of_date": "2026-07-01",
                    "generate_ai_summary": True,
                },
            )
        )
        forecast_id = assert_identifier(forecast["forecast_id"], label="forecast_id")
        assert forecast["status"] == "completed"
        assert forecast["series_errors"] == []
        assert forecast["disclaimer"]

        summary = forecast["summary"]
        # Forecast plus insufficient-data must account for every series: a
        # series that could not be modelled is reported, never dropped.
        assert (
            summary["forecast_count"] + summary["insufficient_data_count"]
            == summary["series_count"]
        )
        assert summary["forecast_count"] > 0

        listed = ok(client.get(f"{V1}/inventory/forecasts", params={"limit": 5}))
        assert_pagination(listed, items_key="forecasts")
        assert forecast_id in {row["forecast_id"] for row in listed["forecasts"]}

        items = ok(
            client.get(f"{V1}/inventory/forecasts/{forecast_id}/items", params={"limit": 10})
        )
        assert_pagination(items, items_key="items")
        for item in items["items"]:
            assert item["output_origin"] in VALID_ORIGINS
            assert item["model"], "every forecast must say which model produced it"
            assert item["selection_basis"], "and why that model was chosen"

        first = items["items"][0]
        detail = ok(
            client.get(
                f"{V1}/inventory/forecasts/{forecast_id}/item",
                params={"material": first["material"], "plant": first["plant"]},
            )
        )
        assert detail["material"] == first["material"]
        assert detail["forecast"], "no forecast points were returned"
        assert detail["model"] == first["model"]

        # The interval has to contain the point it is an interval around.
        for point in detail["forecast"]:
            assert point["lower"] <= point["demand"] <= point["upper"]
            assert point["output_origin"] == "forecast"

        # A reorder recommendation must not contradict the same engine's own
        # shortage prediction. Ordering *after* the shortage lands is only
        # defensible in one situation: quantity is already on order and simply
        # arrives too late, so the action is to expedite that delivery rather
        # than raise another. The pair is what has to hold - a late reorder date
        # with no expedite flag next to it is the module 7 bug returning.
        for item in items["items"]:
            shortage = item["predicted_shortage_date"]
            reorder = item["recommended_reorder_date"]
            if shortage and reorder and reorder > shortage:
                assert item["expedite_recommended"] is True, (
                    f"{item['material']}/{item['plant']} is told to reorder on {reorder}, "
                    f"after its own predicted shortage on {shortage}, with no expedite flag"
                )

        download(
            client.get(
                f"{V1}/inventory/forecasts/{forecast_id}/export", params={"format": "xlsx"}
            ),
            expect_extension="xlsx",
        )


# ===========================================================================
# Module 8 - SAP Test Case Generator
# ===========================================================================
class TestTestCaseGeneratorWorkflow:
    """Generate a suite, edit a case, approve it, run it, export the pack."""

    def test_the_whole_journey(self, client):
        suite = ok(
            client.post(
                f"{V1}/test-cases/generate",
                json={
                    "context": PROCESS_CONTEXT,
                    "test_case_count": 8,
                    "test_types": ["sit", "uat", "negative", "integration"],
                    "use_ai": True,
                },
            )
        )
        suite_id = assert_identifier(suite["suite_id"], label="suite_id")
        cases = suite["test_cases"]
        assert len(cases) == 8

        # The skeleton is deterministic: identifiers, coverage and priorities.
        identifiers = [case["test_case_id"] for case in cases]
        assert len(set(identifiers)) == len(identifiers), "an identifier was reused"
        # Each type numbers its own cases from 1 upwards without gaps.
        by_type: dict[str, list[str]] = {}
        for case in cases:
            by_type.setdefault(case["test_type"], []).append(case["test_case_id"])
        for type_name, ids in by_type.items():
            suffixes = sorted(int(identifier.rsplit("-", 1)[-1]) for identifier in ids)
            assert suffixes == list(range(1, len(ids) + 1)), (
                f"{type_name} identifiers are not numbered 1..n: {ids}"
            )

        for case in cases:
            assert case["output_origin"] in VALID_ORIGINS
            assert case["steps"], "a test case with no steps is not a test case"
            assert [step["step_number"] for step in case["steps"]] == list(
                range(1, len(case["steps"]) + 1)
            )

        listed = ok(client.get(f"{V1}/test-cases/suites", params={"limit": 5}))
        assert_pagination(listed, items_key="suites")
        assert suite_id in {row["suite_id"] for row in listed["suites"]}

        case_id = cases[0]["id"]

        # Approve, then rewrite the script. The approval must not survive a
        # change to the thing that was approved.
        approved = ok(
            client.post(
                f"{V1}/test-cases/{case_id}/approve", json={"approved_by": "Ingrid Halvorsen"}
            )
        )
        assert approved["approved_by"] == "Ingrid Halvorsen"

        edited = ok(
            client.put(
                f"{V1}/test-cases/{case_id}",
                json={
                    "steps": [
                        {"step_number": 1, "action": "Open ME21N", "expected_result": "Screen opens"},
                        {"step_number": 2, "action": "Enter vendor", "expected_result": "Vendor accepted"},
                    ]
                },
            )
        )
        assert edited["approved_by"] is None, (
            "a script edit must clear the approval the old script had earned"
        )

        # Record an execution, then rewrite the script again: the verdict is
        # kept, flagged and counted separately rather than silently retained.
        executed = ok(
            client.post(
                f"{V1}/test-cases/{case_id}/execution",
                json={"execution_result": "failed", "actual_result": "The release strategy did not fire."},
            )
        )
        assert executed["execution_result"] == "failed"
        assert executed["execution_is_stale"] is False

        rewritten = ok(
            client.put(
                f"{V1}/test-cases/{case_id}",
                json={"steps": [{"step_number": 1, "action": "Open ME23N", "expected_result": "Screen opens"}]},
            )
        )
        assert rewritten["execution_is_stale"] is True, (
            "a verdict reached against a script that no longer exists must be flagged"
        )

        refreshed = ok(client.get(f"{V1}/test-cases/suites/{suite_id}"))
        assert refreshed["summary"]["stale_execution_count"] >= 1

        download(
            client.get(f"{V1}/test-cases/suites/{suite_id}/export", params={"format": "xlsx"}),
            expect_extension="xlsx",
        )
        download(
            client.get(f"{V1}/test-cases/suites/{suite_id}/export", params={"format": "pdf"}),
            expect_extension="pdf",
        )


# ===========================================================================
# Module 9 - SAP Blueprint Generator
# ===========================================================================
class TestBlueprintGeneratorWorkflow:
    """Generate a blueprint, edit a section, approve it, version it, export it."""

    def test_the_whole_journey(self, client):
        blueprint = ok(
            client.post(
                f"{V1}/blueprints/generate",
                json={"project": PROJECT, "blueprint_name": "Nordwind e2e", "use_ai": True},
            )
        )
        blueprint_id = assert_identifier(blueprint["blueprint_id"], label="blueprint_id")
        sections = blueprint["sections"]
        assert len(sections) >= 25

        for section in sections:
            assert section["section_id"], "every section needs a stable identifier"
            assert section["title"]
            for item in section["items"]:
                assert item["output_origin"] in VALID_ORIGINS

        listed = ok(client.get(f"{V1}/blueprints", params={"limit": 5}))
        assert_pagination(listed, items_key="blueprints")
        assert blueprint_id in {row["blueprint_id"] for row in listed["blueprints"]}

        # Approve the executive summary, then rewrite the scope it describes.
        ok(
            client.post(
                f"{V1}/blueprints/{blueprint_id}/sections/executive_summary/approve",
                json={"approved_by": "Ingrid Halvorsen"},
            )
        )
        ok(
            client.put(
                f"{V1}/blueprints/{blueprint_id}/sections/scope",
                json={"narrative": "The scope was renegotiated after the summary was approved."},
            )
        )
        summary = ok(
            client.get(f"{V1}/blueprints/{blueprint_id}/sections/executive_summary")
        )
        assert summary["approved_by"] == "Ingrid Halvorsen", "the approval really was given"
        assert summary["approval_is_stale"] is True, (
            "an approval given to a description of something that has since changed must say so"
        )

        detail = ok(client.get(f"{V1}/blueprints/{blueprint_id}"))
        assert detail["summary"]["stale_approved_count"] >= 1

        # Versions: saved, listed and comparable.
        version = ok(
            client.post(f"{V1}/blueprints/{blueprint_id}/versions", json={}), expected_status=201
        )
        assert version["version_number"] >= 1
        versions = ok(client.get(f"{V1}/blueprints/{blueprint_id}/versions"))
        assert versions["versions"]

        download(
            client.get(f"{V1}/blueprints/{blueprint_id}/export", params={"format": "markdown"}),
            expect_extension="md",
        )
        download(
            client.get(f"{V1}/blueprints/{blueprint_id}/export", params={"format": "docx"}),
            expect_extension="docx",
        )


# ===========================================================================
# Module 10 - SAP Interview Coach
# ===========================================================================
class TestInterviewCoachWorkflow:
    """Start a session, answer every question, complete it, read the dashboard."""

    def test_the_whole_journey(self, client):
        catalogue = ok(client.get(f"{V1}/interviews/catalog"))
        assert catalogue["tracks"] and catalogue["modes"]

        # A question browsed before it is answered must not carry its answer key.
        browsed = ok(client.get(f"{V1}/interviews/questions", params={"limit": 5}))
        assert_pagination(browsed, items_key="questions")
        for question in browsed["questions"]:
            rendered = json.dumps(question).lower()
            assert "expected_concepts" not in rendered
            assert "reference_answer" not in rendered

        session = ok(
            client.post(
                f"{V1}/interviews/start",
                json={
                    "tracks": ["sap_mm"],
                    "mode": "practice",
                    "question_count": 3,
                    "seed": 20260801,
                    "candidate_name": "E2E candidate",
                },
            ),
            expected_status=201,
        )
        session_id = assert_identifier(session["session_id"], label="session_id")
        assert session["status"] == "in_progress"
        assert session["summary"]["question_count"] == 3

        served = []
        answer_text = (
            "Goods receipt and invoice receipt are matched through the GR/IR clearing account. "
            "The purchase order, the goods receipt and the supplier invoice form the three-way "
            "match, so the company only settles invoices for goods it actually received. An "
            "aged GR/IR balance means one of the three documents is missing."
        )
        for index in range(3):
            current = session["next_question"] if index == 0 else served[-1]["next_question"]
            assert current is not None, "the session ran out of questions early"
            served.append(
                ok(
                    client.post(
                        f"{V1}/interviews/{session_id}/answer",
                        json={"answer_text": answer_text, "seconds_spent": 90},
                    )
                )
            )
            answer = served[-1]["answer"]
            score = answer["score"]
            assert 0 <= score["overall_score"] <= 100
            assert answer["scoring_is_stale"] is False

            # A dimension the question does not test is reported as not
            # applicable rather than scored zero, and the weights of the
            # dimensions that do apply must still sum to one.
            applicable = [d for d in score["dimensions"] if d["applicable"]]
            assert applicable, "no dimension applied to the answer"
            assert sum(d["weight"] for d in applicable) == pytest.approx(1.0, abs=0.01)
            for dimension in score["dimensions"]:
                assert dimension["explanation"], (
                    "a dimension score with no explanation cannot be argued with"
                )

            # The answer key only ever travels with a submitted answer.
            assert answer["answer_key"]["expected_concepts"]

        # The question served next must actually change - asserting on counts
        # alone cannot see a response that re-serves the question just answered.
        asked = [item["answer"]["question"]["question_id"] for item in served]
        assert len(set(asked)) == 3, f"the same question was served twice: {asked}"

        completed = ok(client.post(f"{V1}/interviews/{session_id}/complete", json={}))
        assert completed["status"] == "completed"
        assert completed["summary"]["answered_count"] == 3
        assert completed["summary"]["pending_count"] == 0

        listed = ok(client.get(f"{V1}/interviews/sessions", params={"limit": 5}))
        assert_pagination(listed, items_key="sessions")
        assert session_id in {row["session_id"] for row in listed["sessions"]}

        dashboard = ok(client.get(f"{V1}/interviews/performance"))
        assert dashboard["session_count"] >= 1
        assert dashboard["stale_score_count"] == 0

        # Every study-plan item must be derived from the data it describes,
        # not from the branch that produced it. A topic recommended for study
        # while its own average disproves the sentence next to it is the
        # module 10 bug returning.
        by_topic = {row["topic"]: row for row in dashboard["by_topic"]}
        for item in dashboard["study_plan"]:
            assert item["reason"], "a study plan item with no reason is not actionable"
            assert item["actions"], "a study plan item with no action is not a plan"
            topic = by_topic.get(item["topic"])
            if topic is not None and topic["average_score"] is not None:
                assert item["average_score"] == pytest.approx(topic["average_score"]), (
                    "the plan quotes a different average from the dashboard it sits on"
                )

        # Download the transcript in all four formats.
        workbook_bytes = download(
            client.get(f"{V1}/interviews/{session_id}/export", params={"format": "xlsx"}),
            expect_extension="xlsx",
        )
        workbook = load_workbook(io.BytesIO(workbook_bytes))
        assert "Answers" in workbook.sheetnames
        assert "Study Plan" in workbook.sheetnames

        download(
            client.get(f"{V1}/interviews/{session_id}/export", params={"format": "csv"}),
            expect_extension="csv",
        )
        transcript = download(
            client.get(f"{V1}/interviews/{session_id}/export", params={"format": "pdf"}),
            expect_extension="pdf",
        )
        assert transcript.startswith(b"%PDF-"), "the PDF export is not a PDF"

        report = json.loads(
            download(
                client.get(f"{V1}/interviews/{session_id}/export", params={"format": "json"}),
                expect_extension="json",
            )
        )
        assert report["session"]["session_id"] == session_id
        assert len(report["answers"]) == 3
        assert "PRACTICE FEEDBACK" in report["disclaimer"], (
            "a report of somebody's answers must not read as a formal assessment"
        )
        # The file and the API must agree about the marks. A report is a second
        # serialisation of the same verdict, and a second chance to disagree.
        assert [item["score"]["overall_score"] for item in report["answers"]] == [
            item["answer"]["score"]["overall_score"] for item in served
        ]
        assert report["summary"]["average_score"] == completed["summary"]["average_score"]

        # An unknown session is a clean 404, not a stack trace.
        error = failure(
            client.get(f"{V1}/interviews/does-not-exist/export"), expected_status=404
        )
        assert error["code"] == "not_found"
        assert "Traceback" not in error["message"]
