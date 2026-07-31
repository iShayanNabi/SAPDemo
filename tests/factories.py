"""Helpers for building small, readable datasets in unit tests.

A rule test should read like a sentence: "two orders on the same supplier with
the same value three days apart produce one duplicate finding". These factories
supply sensible defaults so a test only writes the fields it cares about.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, timedelta
from typing import Any

import pandas as pd
from openpyxl import Workbook

from app.modules.po_risk.field_definitions import CANONICAL_FIELDS
from app.modules.po_risk.normalizer import normalize_dataframe
from app.modules.po_risk.rules.base import RuleContext
from app.modules.po_risk.thresholds import PoRiskConfig, get_rule_config

BASE_DATE = date(2025, 5, 12)

DEFAULT_ROW: dict[str, Any] = {
    "po_number": "4500000001",
    "po_item": "00010",
    "supplier_id": "0000100001",
    "supplier_name": "Nordwind Industrie GmbH",
    "material": "MAT-100000",
    "material_description": "Stainless bearing type 1",
    "material_group": "MG10",
    "company_code": "1000",
    "purchasing_org": "1000",
    "purchasing_group": "P01",
    "plant": "1010",
    "quantity": 10.0,
    "unit_of_measure": "PC",
    "unit_price": 100.0,
    "currency": "EUR",
    "total_value": 1000.0,
    "order_date": BASE_DATE,
    "requested_delivery_date": BASE_DATE + timedelta(days=14),
    "actual_delivery_date": BASE_DATE + timedelta(days=15),
    "contract_number": None,
    "payment_terms": "NT30",
    "approval_status": "Approved",
    "created_by": "BUYER01",
    "changed_by": "BUYER01",
    "change_count": 1,
}


def make_row(**overrides: Any) -> dict[str, Any]:
    """Build one canonical row, applying ``overrides`` on top of the defaults.

    ``total_value`` is recomputed automatically unless it is overridden.
    """
    row = {**DEFAULT_ROW, **overrides}
    if "total_value" not in overrides and row["quantity"] is not None and row["unit_price"] is not None:
        row["total_value"] = round(float(row["quantity"]) * float(row["unit_price"]), 2)
    return row


def make_rows(count: int, **overrides: Any) -> list[dict[str, Any]]:
    """Build ``count`` rows with incrementing purchase order numbers."""
    rows = []
    for index in range(count):
        row = make_row(**overrides)
        row["po_number"] = str(4500000001 + index)
        rows.append(row)
    return rows


def make_frame(rows: list[dict[str, Any]], config: PoRiskConfig | None = None) -> pd.DataFrame:
    """Run rows through the real normaliser to get a canonical frame."""
    config = config or get_rule_config()
    raw = pd.DataFrame([{key: row.get(key) for key in CANONICAL_FIELDS} for row in rows])
    mapping = {name: name for name in CANONICAL_FIELDS}
    return normalize_dataframe(raw, mapping, config).frame


def make_context(rows: list[dict[str, Any]], config: PoRiskConfig | None = None) -> RuleContext:
    """Build a :class:`RuleContext` from raw rows."""
    config = config or get_rule_config()
    return RuleContext(frame=make_frame(rows, config), config=config)


# ---------------------------------------------------------------------------
# File builders for the upload/API tests
# ---------------------------------------------------------------------------


def rows_to_csv(rows: list[dict[str, Any]], headers: dict[str, str] | None = None) -> bytes:
    """Serialise rows as CSV, optionally renaming the headers."""
    headers = headers or {}
    field_names = [headers.get(name, name) for name in CANONICAL_FIELDS]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=field_names, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {headers.get(name, name): _as_text(row.get(name)) for name in CANONICAL_FIELDS}
        )
    return buffer.getvalue().encode("utf-8")


def rows_to_xlsx(rows: list[dict[str, Any]], headers: dict[str, str] | None = None) -> bytes:
    """Serialise rows as a single sheet XLSX workbook."""
    headers = headers or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Purchase Orders"
    sheet.append([headers.get(name, name) for name in CANONICAL_FIELDS])
    for row in rows:
        sheet.append([_as_text(row.get(name)) for name in CANONICAL_FIELDS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def rows_to_json(rows: list[dict[str, Any]], wrap: bool = True) -> bytes:
    """Serialise rows as JSON, either wrapped in ``records`` or as a bare list."""
    payload = [{name: _as_text(row.get(name)) for name in CANONICAL_FIELDS} for row in rows]
    document: Any = {"records": payload} if wrap else payload
    return json.dumps(document, indent=2).encode("utf-8")


def _as_text(value: Any) -> Any:
    """Render a value the way a spreadsheet export would."""
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return value


# ---------------------------------------------------------------------------
# Spend Analytics factories
# ---------------------------------------------------------------------------

SPEND_BASE_DATE = date(2025, 3, 10)

DEFAULT_SPEND_ROW: dict[str, Any] = {
    **DEFAULT_ROW,
    "po_number": "5500000001",
    "supplier_id": "0000200001",
    "supplier_name": "Nordwind Industrie GmbH",
    "material": "SPM-100000",
    "material_description": "Stainless bearing type 1",
    "material_group": "MG10",
    "contract_number": "4600000123",
    "order_date": SPEND_BASE_DATE - timedelta(days=5),
    "requested_delivery_date": SPEND_BASE_DATE + timedelta(days=14),
    "actual_delivery_date": SPEND_BASE_DATE + timedelta(days=15),
    "transaction_date": SPEND_BASE_DATE,
    "category": "Components",
    "subcategory": "Mechanical",
    "contract_status": "Contracted",
    "preferred_supplier_status": "Preferred",
    "baseline_price": 100.0,
    "current_price": 100.0,
    "payment_status": "Paid",
}


def make_spend_row(**overrides: Any) -> dict[str, Any]:
    """Build one canonical spend row on top of a healthy default.

    The default is contracted, preferred and priced exactly at its baseline, so
    a test only has to state the one thing it is exercising.
    """
    row = {**DEFAULT_SPEND_ROW, **overrides}
    if "total_value" not in overrides and row["quantity"] is not None and row["unit_price"] is not None:
        row["total_value"] = round(float(row["quantity"]) * float(row["unit_price"]), 2)
    if "current_price" not in overrides and "unit_price" in overrides:
        row["current_price"] = row["unit_price"]
    if "baseline_price" not in overrides and "unit_price" in overrides:
        row["baseline_price"] = row["unit_price"]
    return row


def make_spend_rows(count: int, **overrides: Any) -> list[dict[str, Any]]:
    """Build ``count`` spend rows with incrementing purchase order numbers."""
    rows = []
    for index in range(count):
        row = make_spend_row(**overrides)
        row["po_number"] = str(5500000001 + index)
        rows.append(row)
    return rows


def make_spend_frame(rows: list[dict[str, Any]], config: Any = None) -> "pd.DataFrame":
    """Run spend rows through the real normaliser to get a canonical frame."""
    from app.modules.spend.field_definitions import CANONICAL_FIELDS as SPEND_FIELDS
    from app.modules.spend.normalizer import normalize_spend_dataframe
    from app.modules.spend.thresholds import get_spend_config

    config = config or get_spend_config()
    raw = pd.DataFrame([{key: row.get(key) for key in SPEND_FIELDS} for row in rows])
    mapping = {name: name for name in SPEND_FIELDS}
    return normalize_spend_dataframe(raw, mapping, config).frame


def spend_rows_to_csv(
    rows: list[dict[str, Any]], headers: dict[str, str] | None = None
) -> bytes:
    """Serialise spend rows as CSV, optionally renaming the headers."""
    from app.modules.spend.field_definitions import CANONICAL_FIELDS as SPEND_FIELDS

    headers = headers or {}
    field_names = [headers.get(name, name) for name in SPEND_FIELDS]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=field_names, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({headers.get(n, n): _as_text(row.get(n)) for n in SPEND_FIELDS})
    return buffer.getvalue().encode("utf-8")


def spend_rows_to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """Serialise spend rows as an XLSX workbook."""
    from app.modules.spend.field_definitions import CANONICAL_FIELDS as SPEND_FIELDS

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Spend"
    sheet.append(list(SPEND_FIELDS))
    for row in rows:
        sheet.append([_as_text(row.get(name)) for name in SPEND_FIELDS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Supplier Recommendation factories
# ---------------------------------------------------------------------------

#: A healthy, well-rounded supplier. A test overrides only what it exercises.
DEFAULT_SUPPLIER: dict[str, Any] = {
    "supplier_id": "0000300001",
    "supplier_name": "Nordwind Industrie GmbH",
    "materials_supplied": ["MAT-1000"],
    "plants_served": ["1010"],
    "regions_served": ["EU"],
    "unit_price": 100.0,
    "currency": "EUR",
    "lead_time_days": 14,
    "available_capacity": 1000.0,
    "on_time_delivery_rate": 95.0,
    "quality_score": 90.0,
    "defect_rate": 1.5,
    "risk_score": 25.0,
    "esg_score": 80.0,
    "contract_status": "Active",
    "contract_expiration": date(2028, 1, 1),
    "payment_terms": "NT30",
    "historical_order_count": 40,
    "historical_spend": 400000.0,
}


def make_supplier(**overrides: Any):
    """Build one :class:`NormalizedSupplier` on top of a healthy default."""
    from app.modules.supplier_reco.normalizer import NormalizedSupplier
    from app.modules.supplier_reco.thresholds import get_supplier_reco_config

    data = {**DEFAULT_SUPPLIER, **overrides}
    config = get_supplier_reco_config()
    rate = config.conversion_rate(data["currency"])
    price = data["unit_price"]
    spend = data["historical_spend"]
    return NormalizedSupplier(
        row_number=overrides.get("row_number", 2),
        supplier_id=str(data["supplier_id"]),
        supplier_name=data["supplier_name"],
        materials_supplied=list(data["materials_supplied"]),
        plants_served=list(data["plants_served"]),
        regions_served=list(data["regions_served"]),
        unit_price=price,
        currency=data["currency"],
        unit_price_base=None if price is None else round(price * rate, 4),
        lead_time_days=data["lead_time_days"],
        available_capacity=data["available_capacity"],
        on_time_delivery_rate=data["on_time_delivery_rate"],
        quality_score=data["quality_score"],
        defect_rate=data["defect_rate"],
        risk_score=data["risk_score"],
        esg_score=data["esg_score"],
        contract_status=data["contract_status"],
        contract_expiration=data["contract_expiration"],
        payment_terms=data["payment_terms"],
        historical_order_count=data["historical_order_count"],
        historical_spend=spend,
        historical_spend_base=None if spend is None else round(spend * rate, 2),
    )


def make_suppliers(count: int, **overrides: Any) -> list[Any]:
    """Build ``count`` suppliers with incrementing supplier ids."""
    suppliers = []
    for index in range(count):
        data = dict(overrides)
        data.setdefault("supplier_id", f"{300001 + index:010d}")
        if "supplier_id" not in overrides:
            data["supplier_id"] = f"{300001 + index:010d}"
        suppliers.append(make_supplier(**data))
    return suppliers


def make_requirement(**overrides: Any):
    """Build a :class:`Requirement` with sensible defaults for the canonical material."""
    from app.modules.supplier_reco.requirement import Requirement

    defaults: dict[str, Any] = {
        "material": "MAT-1000",
        "quantity": 100.0,
        "plant": "1010",
        "preferred_region": "EU",
        "required_delivery_date": date(2026, 10, 1),
        "order_date": date(2026, 8, 1),
        "target_price": 120.0,
        "currency": "EUR",
        "risk_tolerance": "medium",
    }
    defaults.update(overrides)
    return Requirement(**defaults)


def supplier_rows_to_csv(rows: list[dict[str, Any]], headers: dict[str, str] | None = None) -> bytes:
    """Serialise supplier rows as CSV, optionally renaming the headers.

    List fields (materials/plants/regions) are joined with ``;``.
    """
    from app.modules.supplier_reco.field_definitions import CANONICAL_FIELDS as SUPPLIER_FIELDS

    headers = headers or {}
    field_names = [headers.get(name, name) for name in SUPPLIER_FIELDS]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=field_names, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {headers.get(name, name): _supplier_cell(row.get(name)) for name in SUPPLIER_FIELDS}
        )
    return buffer.getvalue().encode("utf-8")


def supplier_rows_to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """Serialise supplier rows as an XLSX workbook."""
    from app.modules.supplier_reco.field_definitions import CANONICAL_FIELDS as SUPPLIER_FIELDS

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Suppliers"
    sheet.append(list(SUPPLIER_FIELDS))
    for row in rows:
        sheet.append([_supplier_cell(row.get(name)) for name in SUPPLIER_FIELDS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def supplier_row(**overrides: Any) -> dict[str, Any]:
    """Build a raw canonical supplier row dict (lists as Python lists) for file builders."""
    return {**DEFAULT_SUPPLIER, **overrides}


def _supplier_cell(value: Any) -> Any:
    """Render a supplier cell the way a file export would."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ";".join(str(item) for item in value)
    if isinstance(value, date):
        return value.isoformat()
    return value
