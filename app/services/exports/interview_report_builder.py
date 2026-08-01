"""Report builders for the SAP Interview Coach.

Four formats, all built from one payload:

* **XLSX** - five sheets: Summary, Answers, Dimension Scores, Concept Coverage
  and Study Plan.
* **CSV** - one row per answer, which is what goes into a tracker.
* **JSON** - the whole session, for a downstream system.
* **PDF** - the transcript, one answer per page, read top to bottom.

The PDF uses the same dependency-free text writer the Test Case Generator uses,
for the same reason: this document is a column of text somebody reads, not a
laid-out report, and that is exactly what ``build_text_pdf`` produces.

Three things this report has to get right, all of them about honesty rather than
completeness.

**The score and the prose have different origins, and the file has to say so.**
The marks come from the rubric - deterministic, reproducible, identical with a
real model, with the mock, or with no AI at all. The coaching note, the sample
answer and the strengths are *prose*, drafted by whatever provider was active.
Every sheet that carries prose carries its origin next to it, because a candidate
who cannot tell which is which will treat a mock paragraph as a verdict.

**A concept that was credited has to say what credited it.** The rubric matches
keywords, and keyword matching has a ceiling. Printing the matched keyword and
the excerpt next to each concept is what lets somebody disagree with a specific
decision instead of with a number.

**A stale score is exported as stale.** A mark computed against a rubric that has
since been retuned is kept, flagged and counted separately in the API; a report
that quietly dropped the flag would be the one place that lie survives.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.logging import get_logger
from app.services.documents.pdf_writer import build_text_pdf
from app.services.exports.styling import (
    SECTION_FONT,
    TITLE_FONT,
    WRAP,
    cell_value,
    humanise,
    scalar,
    write_header,
    write_key_values,
    write_paragraph,
)
from app.services.exports.workbook import workbook_to_bytes

logger = get_logger(__name__)

#: (payload key, column label) for the per-answer table. Nested values are
#: resolved by :func:`_answer_row` before this is applied.
ANSWER_COLUMNS: tuple[tuple[str, str], ...] = (
    ("position", "#"),
    ("question_id", "Question ID"),
    ("track", "Track"),
    ("topic", "Topic"),
    ("difficulty", "Difficulty"),
    ("status", "Status"),
    ("question", "Question"),
    ("answer_text", "Candidate Answer"),
    ("overall_score", "Overall Score"),
    ("overall_band", "Band"),
    ("passed", "Passed"),
    ("concepts_matched", "Concepts Matched"),
    ("concepts_expected", "Concepts Expected"),
    ("seconds_spent", "Seconds Spent"),
    ("within_time_limit", "Within Time Limit"),
    ("scoring_is_stale", "Score Is Stale"),
    ("strengths", "Strengths"),
    ("missing_concepts", "Missing Concepts"),
    ("incorrect_statements", "Incorrect Statements"),
    ("coaching_note", "Coaching Note"),
    ("improved_sample_answer", "Improved Sample Answer"),
    ("topics_to_study", "Topics To Study"),
    ("feedback_origin", "Feedback Origin"),
    ("score_origin", "Score Origin"),
)


def interview_disclaimer() -> str:
    """The disclaimer printed on every interview export.

    Two claims, and the second is the one a candidate needs. The scores come
    from a keyword rubric, which is reproducible and *limited*: it can miss a
    correct answer phrased in words the rubric does not know. Saying so on the
    file is the difference between a practice tool and a pretend qualification.
    """
    return (
        "Scores in this report are produced by a deterministic keyword rubric from the answer "
        "text only. They are identical with a real AI provider, with the built-in mock, or with "
        "AI switched off entirely - no model contributes to any mark. THIS IS PRACTICE FEEDBACK, "
        "NOT AN ASSESSMENT: keyword matching has a ceiling and can miss a correct answer phrased "
        "in words the rubric does not hold, which is why the keyword that credited each concept "
        "is printed next to it. Coaching prose, sample answers and study suggestions are written "
        "by the configured AI provider (or the built-in templates) and are labelled with their "
        "origin; they never change a score. Time spent is recorded and reported and never affects "
        "a mark. This application is not connected to any SAP system and no result here is an SAP "
        "certification, qualification or endorsement."
    )


# ---------------------------------------------------------------------------
# JSON / CSV
# ---------------------------------------------------------------------------


def build_interview_json_report(payload: dict[str, Any]) -> bytes:
    """Serialise the whole session payload as JSON."""
    document = {
        "report_type": "sap_interview_session",
        "generated_at": datetime.now(UTC).isoformat(),
        "disclaimer": interview_disclaimer(),
        **payload,
    }
    return json.dumps(document, indent=2, ensure_ascii=False, default=str).encode("utf-8")


def build_interview_csv_report(answers: list[dict[str, Any]]) -> bytes:
    """Serialise one row per answer as CSV (UTF-8 with BOM, so Excel opens it)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([label for _, label in ANSWER_COLUMNS])
    for answer in answers:
        row = _answer_row(answer)
        writer.writerow([cell_value(row, key) for key, _ in ANSWER_COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def build_interview_xlsx_report(payload: dict[str, Any]) -> bytes:
    """Build the multi-sheet interview workbook."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    answers = payload.get("answers", []) or []

    _summary_sheet(workbook.create_sheet("Summary"), payload)
    _answers_sheet(workbook.create_sheet("Answers"), answers)
    _dimensions_sheet(workbook.create_sheet("Dimension Scores"), answers)
    _concepts_sheet(workbook.create_sheet("Concept Coverage"), answers)
    _study_sheet(workbook.create_sheet("Study Plan"), payload)

    payload_bytes = workbook_to_bytes(workbook)
    logger.info("Built interview session XLSX: %d answers", len(answers))
    return payload_bytes


def _summary_sheet(sheet: Worksheet, payload: dict[str, Any]) -> None:
    session = payload.get("session", {}) or {}
    summary = payload.get("summary", {}) or {}

    sheet["A1"] = "SAP Interview Coach - Session Report"
    sheet["A1"].font = TITLE_FONT
    row = write_paragraph(sheet, interview_disclaimer(), 2, height=7)

    row += 1
    sheet.cell(row=row, column=1, value="Session").font = SECTION_FONT
    row = write_key_values(
        sheet,
        [
            ("Session ID", session.get("session_id")),
            ("Name", session.get("name")),
            ("Candidate", session.get("candidate_name")),
            ("Status", session.get("status")),
            ("Tracks", session.get("tracks")),
            ("Mode", session.get("mode")),
            ("Difficulties", session.get("difficulties")),
            ("Topics", session.get("topics")),
            ("Questions requested", session.get("requested_question_count")),
            ("Selection seed", session.get("seed")),
            ("Time limit (seconds)", session.get("time_limit_seconds")),
            ("Started", session.get("started_at")),
            ("Completed", session.get("completed_at")),
            ("Question bank version", session.get("question_bank_version")),
            ("Configuration version", session.get("config_version")),
            ("Engine version", session.get("engine_version")),
        ],
        row + 1,
    )

    row += 1
    sheet.cell(row=row, column=1, value="Result (rule-based)").font = SECTION_FONT
    row = write_key_values(
        sheet,
        [
            (humanise(key), value)
            for key, value in summary.items()
            if key not in {"average_by_dimension", "average_by_topic", "average_by_difficulty"}
        ],
        row + 1,
    )

    for key, label in (
        ("average_by_dimension", "Average by dimension"),
        ("average_by_topic", "Average by topic"),
        ("average_by_difficulty", "Average by difficulty"),
    ):
        breakdown = summary.get(key) or {}
        if not breakdown:
            continue
        row += 1
        sheet.cell(row=row, column=1, value=label).font = SECTION_FONT
        row = write_key_values(
            sheet, [(humanise(str(k)), v) for k, v in breakdown.items()], row + 1
        )

    notes = payload.get("notes") or []
    if notes:
        row += 1
        sheet.cell(row=row, column=1, value="Notes").font = SECTION_FONT
        row += 1
        for note in notes:
            sheet.cell(row=row, column=1, value=scalar(note)).alignment = WRAP
            row += 1

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 60


def _answers_sheet(sheet: Worksheet, answers: list[dict[str, Any]]) -> None:
    write_header(sheet, [label for _, label in ANSWER_COLUMNS])
    for index, answer in enumerate(answers, start=2):
        row = _answer_row(answer)
        for column, (key, _label) in enumerate(ANSWER_COLUMNS, start=1):
            cell = sheet.cell(row=index, column=column, value=cell_value(row, key))
            if key in {"question", "answer_text", "coaching_note", "improved_sample_answer"}:
                cell.alignment = WRAP
    sheet.freeze_panes = "B2"
    for letter, width in (
        ("G", 60), ("H", 60), ("Q", 40), ("R", 40), ("S", 40), ("T", 60), ("U", 60), ("V", 34)
    ):
        sheet.column_dimensions[letter].width = width


def _dimensions_sheet(sheet: Worksheet, answers: list[dict[str, Any]]) -> None:
    """One row per answer per dimension, including the ones that did not apply."""
    sheet["A1"] = "Dimension scores (rule-based)"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "A dimension a question does not test is reported as not applicable rather "
        "than scored zero, and its weight is shared among the dimensions that do "
        "apply. The applicable weights sum to 1."
    )
    sheet["A2"].alignment = WRAP
    write_header(
        sheet,
        [
            "#",
            "Question ID",
            "Topic",
            "Dimension",
            "Applicable",
            "Score",
            "Weight",
            "Band",
            "Concepts Expected",
            "Concepts Matched",
            "Explanation",
        ],
        row=4,
    )

    row = 5
    for answer in answers:
        question = answer.get("question", {}) or {}
        score = answer.get("score") or {}
        for dimension in score.get("dimensions", []) or []:
            values = [
                answer.get("position"),
                question.get("question_id"),
                question.get("topic"),
                dimension.get("dimension"),
                dimension.get("applicable"),
                dimension.get("score"),
                dimension.get("weight"),
                dimension.get("band"),
                dimension.get("concepts_expected"),
                dimension.get("concepts_matched"),
                dimension.get("explanation"),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=scalar(value))
            row += 1

    sheet.freeze_panes = "A5"
    sheet.column_dimensions["K"].width = 70


def _concepts_sheet(sheet: Worksheet, answers: list[dict[str, Any]]) -> None:
    """Which concepts were credited, and what credited them."""
    sheet["A1"] = "Concept coverage (rule-based)"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "The keyword that credited each concept is printed so a decision can be "
        "disagreed with specifically. A concept marked as not covered may still "
        "have been answered correctly in words the rubric does not hold."
    )
    sheet["A2"].alignment = WRAP
    write_header(
        sheet,
        [
            "#",
            "Question ID",
            "Concept",
            "Dimension",
            "Required",
            "Weight",
            "Covered",
            "Matched Keyword",
            "Excerpt",
            "Negated",
            "Study Hint",
        ],
        row=4,
    )

    row = 5
    for answer in answers:
        question = answer.get("question", {}) or {}
        score = answer.get("score") or {}
        for match in score.get("concept_matches", []) or []:
            values = [
                answer.get("position"),
                question.get("question_id"),
                match.get("label"),
                match.get("dimension"),
                match.get("required"),
                match.get("weight"),
                match.get("matched"),
                match.get("matched_keyword"),
                match.get("excerpt"),
                match.get("negated"),
                match.get("study_hint"),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=scalar(value))
            row += 1

    sheet.freeze_panes = "A5"
    sheet.column_dimensions["C"].width = 48
    sheet.column_dimensions["I"].width = 60
    sheet.column_dimensions["K"].width = 60


def _study_sheet(sheet: Worksheet, payload: dict[str, Any]) -> None:
    summary = payload.get("summary", {}) or {}
    answers = payload.get("answers", []) or []

    sheet["A1"] = "What to study next"
    sheet["A1"].font = TITLE_FONT

    row = 3
    sheet.cell(row=row, column=1, value="Recommended topics (from this session)").font = (
        SECTION_FONT
    )
    row += 1
    for topic in summary.get("recommended_study_topics", []) or []:
        sheet.cell(row=row, column=1, value=scalar(topic))
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Most missed concepts").font = SECTION_FONT
    row += 1
    for concept in summary.get("most_missed_concepts", []) or []:
        sheet.cell(row=row, column=1, value=scalar(concept))
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Strong topics").font = SECTION_FONT
    row = write_key_values(sheet, [("", summary.get("strong_topics"))], row + 1)
    sheet.cell(row=row, column=1, value="Weak topics").font = SECTION_FONT
    row = write_key_values(sheet, [("", summary.get("weak_topics"))], row + 1)

    row += 1
    sheet.cell(row=row, column=1, value="Per-answer study suggestions").font = SECTION_FONT
    row += 1
    write_header(sheet, ["#", "Question ID", "Topics To Study", "Origin"], row)
    row += 1
    for answer in answers:
        feedback = answer.get("feedback") or {}
        topics = feedback.get("topics_to_study") or []
        if not topics:
            continue
        question = answer.get("question", {}) or {}
        for column, value in enumerate(
            [
                answer.get("position"),
                question.get("question_id"),
                " | ".join(str(item) for item in topics),
                feedback.get("output_origin"),
            ],
            start=1,
        ):
            sheet.cell(row=row, column=column, value=scalar(value))
        row += 1

    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["C"].width = 70


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def build_interview_pdf_report(payload: dict[str, Any]) -> bytes:
    """Render the session as a readable transcript.

    One answer per page: the question, what the candidate wrote, the marks with
    their reasons, and the coaching prose with its origin. Built on the same
    text PDF writer the Test Case Generator uses.
    """
    session = payload.get("session", {}) or {}
    title = session.get("name") or "SAP interview session"
    return build_text_pdf(_pdf_text(payload), title=str(title)).content


def _pdf_text(payload: dict[str, Any]) -> str:
    """Lay the session out as plain text, one answer per page."""
    session = payload.get("session", {}) or {}
    summary = payload.get("summary", {}) or {}
    lines: list[str] = []

    lines.append("SAP INTERVIEW COACH - SESSION REPORT")
    lines.append("=" * 70)
    lines.append(f"Session:     {session.get('name', '')}")
    lines.append(f"Candidate:   {session.get('candidate_name') or 'not given'}")
    lines.append(f"Tracks:      {', '.join(session.get('tracks') or [])}")
    lines.append(f"Mode:        {session.get('mode', '')}")
    lines.append(f"Difficulty:  {', '.join(session.get('difficulties') or []) or 'any'}")
    lines.append(f"Status:      {session.get('status', '')}")
    lines.append(f"Generated:   {datetime.now(UTC).isoformat(timespec='seconds')}")
    lines.append("")
    lines.append(
        f"Answered {summary.get('answered_count', 0)} of {summary.get('question_count', 0)}"
        f"   Average {summary.get('average_score')}"
        f"   Band {summary.get('band', '')}"
        f"   Passed {summary.get('passed_count', 0)}"
    )
    if summary.get("pending_count"):
        lines.append(
            f"{summary['pending_count']} question(s) unanswered - not counted as zero."
        )
    if summary.get("stale_score_count"):
        lines.append(
            f"{summary['stale_score_count']} score(s) were computed against an earlier "
            "version of the rubric and are flagged in the transcript."
        )
    if summary.get("over_time_count"):
        lines.append(
            f"{summary['over_time_count']} answer(s) ran over the time limit. Time is "
            "recorded and reported; it never changes a mark."
        )
    lines.append("")
    lines.append("Recommended study topics")
    lines.append("-" * 70)
    for topic in summary.get("recommended_study_topics", []) or ["(none)"]:
        lines.append(f"  - {topic}")
    lines.append("")
    lines.append("DISCLAIMER")
    lines.append("-" * 70)
    lines.extend(_wrap(interview_disclaimer()))

    for answer in payload.get("answers", []) or []:
        lines.append("\f")
        lines.extend(_answer_pages(answer))

    return "\n".join(lines)


def _answer_pages(answer: dict[str, Any]) -> list[str]:
    """One answer laid out as a page of text."""
    question = answer.get("question", {}) or {}
    score = answer.get("score") or {}
    feedback = answer.get("feedback") or {}
    lines: list[str] = []

    lines.append(f"QUESTION {answer.get('position', '')} - {question.get('question_id', '')}")
    lines.append("=" * 70)
    lines.append(
        f"Track {question.get('track', '')}   Topic {question.get('topic', '')}   "
        f"Difficulty {question.get('difficulty', '')}   Status {answer.get('status', '')}"
    )
    lines.append("")
    lines.extend(_wrap(str(question.get("question", ""))))
    lines.append("")
    lines.append("Candidate answer")
    lines.append("-" * 70)
    lines.extend(_wrap(str(answer.get("answer_text") or "(not answered)")))
    lines.append("")

    if not score:
        lines.append("Not answered, so not scored. It is not counted as zero.")
        return lines

    lines.append("Score (rule-based)")
    lines.append("-" * 70)
    lines.append(
        f"  Overall {score.get('overall_score')}   Band {score.get('overall_band', '')}   "
        f"{'PASSED' if score.get('passed') else 'not passed'} "
        f"(pass mark {score.get('pass_score')})"
    )
    if answer.get("scoring_is_stale"):
        lines.append(
            "  *** This score was computed against an earlier version of the rubric. ***"
        )
    lines.append(
        f"  Time {answer.get('seconds_spent')}s"
        + ("" if answer.get("within_time_limit", True) else " - over the limit")
    )
    lines.append("")
    for dimension in score.get("dimensions", []) or []:
        if not dimension.get("applicable"):
            lines.append(
                f"  {str(dimension.get('dimension', '')):<22} not applicable to this question"
            )
            continue
        lines.append(
            f"  {str(dimension.get('dimension', '')):<22} {dimension.get('score')!s:>6}"
            f"   weight {dimension.get('weight')}"
        )
        lines.extend(_wrap(str(dimension.get("explanation", "")), indent="      "))

    lines.append("")
    lines.append("Concepts")
    lines.append("-" * 70)
    for match in score.get("concept_matches", []) or []:
        mark = "[x]" if match.get("matched") else "[ ]"
        credited = f"  credited by \"{match.get('matched_keyword')}\"" if match.get("matched") else ""
        lines.append(f"  {mark} {match.get('label', '')}{credited}")

    for label, key in (
        ("Strengths", "strengths"),
        ("Missing concepts", "missing_concepts"),
        ("Incorrect statements", "incorrect_statements"),
        ("Topics to study", "topics_to_study"),
    ):
        values = feedback.get(key) or []
        if not values:
            continue
        lines.append("")
        lines.append(label)
        lines.append("-" * 70)
        for value in values:
            lines.extend(_wrap(f"- {value}"))

    if feedback.get("coaching_note"):
        lines.append("")
        lines.append(f"Coaching note ({feedback.get('output_origin', 'unknown origin')})")
        lines.append("-" * 70)
        lines.extend(_wrap(str(feedback["coaching_note"])))

    if feedback.get("improved_sample_answer"):
        lines.append("")
        lines.append(f"Improved sample answer ({feedback.get('output_origin', 'unknown origin')})")
        lines.append("-" * 70)
        lines.extend(_wrap(str(feedback["improved_sample_answer"])))

    return lines


def _wrap(text: str, *, width: int = 88, indent: str = "") -> list[str]:
    """Wrap prose to the PDF's column width."""
    import textwrap

    if not text:
        return [""]
    wrapped: list[str] = []
    for paragraph in text.splitlines() or [""]:
        wrapped.extend(
            textwrap.wrap(
                paragraph, width=width, initial_indent=indent, subsequent_indent=indent
            )
            or [indent]
        )
    return wrapped


# ---------------------------------------------------------------------------
# Row flattening
# ---------------------------------------------------------------------------


def _answer_row(answer: dict[str, Any]) -> dict[str, Any]:
    """Flatten one answer into the columns the table and CSV share."""
    question = answer.get("question", {}) or {}
    score = answer.get("score") or {}
    feedback = answer.get("feedback") or {}
    matches = score.get("concept_matches", []) or []

    return {
        "position": answer.get("position"),
        "question_id": question.get("question_id"),
        "track": question.get("track"),
        "topic": question.get("topic"),
        "difficulty": question.get("difficulty"),
        "status": answer.get("status"),
        "question": question.get("question"),
        "answer_text": answer.get("answer_text"),
        "overall_score": score.get("overall_score"),
        "overall_band": score.get("overall_band"),
        "passed": score.get("passed"),
        "concepts_matched": sum(1 for match in matches if match.get("matched")),
        "concepts_expected": len(matches),
        "seconds_spent": answer.get("seconds_spent"),
        "within_time_limit": answer.get("within_time_limit"),
        "scoring_is_stale": answer.get("scoring_is_stale"),
        "strengths": feedback.get("strengths"),
        "missing_concepts": feedback.get("missing_concepts"),
        "incorrect_statements": feedback.get("incorrect_statements"),
        "coaching_note": feedback.get("coaching_note"),
        "improved_sample_answer": feedback.get("improved_sample_answer"),
        "topics_to_study": feedback.get("topics_to_study"),
        "feedback_origin": feedback.get("output_origin"),
        "score_origin": score.get("output_origin"),
    }
