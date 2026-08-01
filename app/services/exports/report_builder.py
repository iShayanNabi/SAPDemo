"""Report exports for the Purchase Order Risk Checker.

Three formats are produced from the same in-memory analysis payload:

* **XLSX** - a multi-sheet workbook (summary, findings, suppliers, rule
  catalogue, data quality) intended for sharing with a purchasing team.
* **CSV** - the findings table only, for import into another tool.
* **JSON** - the complete payload, for a future website or a downstream job.

The workbook contains *results*, not a financial model, so cells hold values
rather than formulas: the numbers come from the deterministic engine and must
not silently change when someone edits a cell.
"""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.logging import get_logger
from app.services.exports.workbook import workbook_to_bytes

logger = get_logger(__name__)

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
LABEL_FONT = Font(name=FONT_NAME, bold=True, size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

SEVERITY_FILLS = {
    "critical": PatternFill("solid", fgColor="F8CBAD"),
    "high": PatternFill("solid", fgColor="FCE4D6"),
    "medium": PatternFill("solid", fgColor="FFF2CC"),
    "low": PatternFill("solid", fgColor="E2EFDA"),
}

FINDING_COLUMNS: tuple[tuple[str, str], ...] = (
    ("finding_id", "Finding ID"),
    ("analysis_id", "Analysis ID"),
    ("po_number", "PO Number"),
    ("po_item", "Item"),
    ("supplier_id", "Supplier ID"),
    ("supplier_name", "Supplier Name"),
    ("risk_category", "Risk Category"),
    ("rule_id", "Rule ID"),
    ("rule_name", "Rule Name"),
    ("severity", "Severity"),
    ("explanation", "Explanation (rule-based)"),
    ("evidence_text", "Supporting Evidence"),
    ("recommended_action", "Recommended Action"),
    ("confidence_score", "Confidence"),
    ("estimated_financial_exposure", "Estimated Exposure"),
    ("exposure_currency", "Currency"),
    ("output_origin", "Output Origin"),
    ("ai_explanation", "AI Explanation (optional)"),
    ("created_at", "Created"),
)

SUPPLIER_COLUMNS: tuple[tuple[str, str], ...] = (
    ("supplier_id", "Supplier ID"),
    ("supplier_name", "Supplier Name"),
    ("spend_base", "Spend"),
    ("spend_share_pct", "Spend Share %"),
    ("line_items", "Line Items"),
    ("findings_count", "Findings"),
    ("critical_count", "Critical"),
    ("high_count", "High"),
    ("medium_count", "Medium"),
    ("low_count", "Low"),
    ("estimated_exposure_base", "Estimated Exposure"),
    ("top_risk_category", "Top Risk Category"),
)


def build_json_report(payload: dict[str, Any]) -> bytes:
    """Serialise the full analysis payload as pretty printed JSON."""
    document = {
        "report_type": "sap_po_risk_analysis",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "disclaimer": _disclaimer(),
        **payload,
    }
    return json.dumps(document, indent=2, ensure_ascii=False, default=str).encode("utf-8")


def build_csv_report(findings: list[dict[str, Any]]) -> bytes:
    """Serialise the findings table as CSV."""
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([label for _, label in FINDING_COLUMNS])
    for finding in findings:
        writer.writerow([_cell_value(finding, key) for key, _ in FINDING_COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


def build_xlsx_report(payload: dict[str, Any]) -> bytes:
    """Build the multi-sheet XLSX report."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    _build_summary_sheet(workbook.create_sheet("Summary"), payload)
    _build_findings_sheet(workbook.create_sheet("Findings"), payload.get("findings", []))
    _build_supplier_sheet(workbook.create_sheet("Supplier Risk"), payload.get("supplier_risk", []))
    _build_methodology_sheet(workbook.create_sheet("Methodology"), payload)
    _build_data_quality_sheet(
        workbook.create_sheet("Data Quality"), payload.get("data_quality_issues", [])
    )

    payload_bytes = workbook_to_bytes(workbook)
    logger.info("Built XLSX report with %d findings", len(payload.get("findings", [])))
    return payload_bytes


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------


def _build_summary_sheet(sheet: Worksheet, payload: dict[str, Any]) -> None:
    analysis = payload.get("analysis", {})
    kpis = payload.get("kpis", {})
    narrative = payload.get("ai_narrative", {}) or {}

    sheet["A1"] = "SAP Purchase Order Risk Analysis"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = _disclaimer()
    sheet["A2"].font = Font(name=FONT_NAME, size=9, italic=True)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A2:D2")
    sheet.row_dimensions[2].height = 45

    rows: list[tuple[str, Any]] = [
        ("Analysis ID", analysis.get("analysis_id")),
        ("Source file", analysis.get("source_filename")),
        ("Created", analysis.get("created_at")),
        ("Status", analysis.get("status")),
        ("Rule config version", analysis.get("config_version")),
        ("Engine version", analysis.get("engine_version")),
        ("", ""),
        ("Line items analysed", kpis.get("record_count")),
        ("Purchase orders", kpis.get("purchase_order_count")),
        ("Suppliers", kpis.get("supplier_count")),
        (f"Total value ({kpis.get('base_currency', 'EUR')})", kpis.get("total_value_base")),
        ("", ""),
        ("Findings (rule-based)", kpis.get("findings_count")),
        ("Critical", (kpis.get("severity_counts") or {}).get("critical")),
        ("High", (kpis.get("severity_counts") or {}).get("high")),
        ("Medium", (kpis.get("severity_counts") or {}).get("medium")),
        ("Low", (kpis.get("severity_counts") or {}).get("low")),
        ("Flagged purchase orders", kpis.get("flagged_purchase_orders")),
        ("Flagged value share %", kpis.get("flagged_value_share_pct")),
        ("Estimated exposure", kpis.get("estimated_exposure_base")),
        ("Risk score (0-100)", kpis.get("risk_score")),
        ("Risk score method", kpis.get("risk_score_method")),
        ("Exposure note", kpis.get("exposure_note")),
    ]

    row_index = 4
    for label, value in rows:
        if label:
            sheet.cell(row=row_index, column=1, value=label).font = LABEL_FONT
            cell = sheet.cell(row=row_index, column=2, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_index += 1

    row_index += 1
    origin = narrative.get("origin") or "not generated"
    sheet.cell(row=row_index, column=1, value="Executive summary").font = LABEL_FONT
    sheet.cell(row=row_index, column=2, value=f"Origin: {origin}").font = BODY_FONT
    row_index += 1
    summary_cell = sheet.cell(
        row=row_index,
        column=2,
        value=narrative.get("summary") or "No AI summary was generated for this analysis.",
    )
    summary_cell.font = BODY_FONT
    summary_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index + 4, end_column=5)

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 60
    for column in ("C", "D", "E"):
        sheet.column_dimensions[column].width = 18


def _build_findings_sheet(sheet: Worksheet, findings: list[dict[str, Any]]) -> None:
    _write_header(sheet, [label for _, label in FINDING_COLUMNS])
    for row_index, finding in enumerate(findings, start=2):
        for column_index, (key, _) in enumerate(FINDING_COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_cell_value(finding, key))
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=key in {"explanation", "evidence_text",
                                                         "recommended_action", "ai_explanation"},
                                       vertical="top")
        severity = str(finding.get("severity", "")).lower()
        if severity in SEVERITY_FILLS:
            sheet.cell(row=row_index, column=10).fill = SEVERITY_FILLS[severity]

    widths = {"Explanation (rule-based)": 60, "Supporting Evidence": 50,
              "Recommended Action": 45, "AI Explanation (optional)": 50,
              "Rule Name": 30, "Supplier Name": 26, "Finding ID": 22, "Analysis ID": 22}
    for index, (_, label) in enumerate(FINDING_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(label, 16)
    sheet.freeze_panes = "A2"
    if findings:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(FINDING_COLUMNS))}{len(findings) + 1}"


def _build_supplier_sheet(sheet: Worksheet, suppliers: list[dict[str, Any]]) -> None:
    _write_header(sheet, [label for _, label in SUPPLIER_COLUMNS])
    for row_index, supplier in enumerate(suppliers, start=2):
        for column_index, (key, _) in enumerate(SUPPLIER_COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=supplier.get(key))
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
    for index, (_, label) in enumerate(SUPPLIER_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = 26 if "Name" in label else 16
    sheet.freeze_panes = "A2"


def _build_methodology_sheet(sheet: Worksheet, payload: dict[str, Any]) -> None:
    sheet["A1"] = "Methodology and rule catalogue"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "Every finding in this report is produced by a deterministic Python rule using the "
        "thresholds below. No AI model takes part in the risk decision. AI text, when present, "
        "is clearly labelled and only restates rule output in business language."
    )
    sheet["A2"].font = Font(name=FONT_NAME, size=9, italic=True)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A2:F2")
    sheet.row_dimensions[2].height = 40

    headers = ["Rule ID", "Rule Name", "Category", "Enabled", "Base Severity", "Confidence",
               "Thresholds", "Findings in this analysis"]
    _write_header(sheet, headers, row=4)

    rules: list[dict[str, Any]] = payload.get("rule_catalogue", [])
    counts = {entry["rule_id"]: entry["count"] for entry in payload.get("kpis", {}).get("rule_counts", [])}
    for row_index, rule in enumerate(rules, start=5):
        values = [
            rule.get("rule_id"),
            rule.get("name"),
            rule.get("category"),
            "yes" if rule.get("enabled") else "no",
            rule.get("base_severity"),
            rule.get("confidence"),
            json.dumps(rule.get("params", {}), ensure_ascii=False),
            counts.get(rule.get("rule_id"), 0),
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=column_index == 7, vertical="top")

    for index, width in enumerate([12, 34, 26, 10, 14, 12, 60, 14], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _build_data_quality_sheet(sheet: Worksheet, issues: list[dict[str, Any]]) -> None:
    headers = ["Field", "Issue Type", "Severity", "Affected Rows", "Sample Rows", "Message"]
    _write_header(sheet, headers)
    for row_index, issue in enumerate(issues, start=2):
        values = [
            issue.get("field"),
            issue.get("issue_type"),
            issue.get("severity"),
            issue.get("affected_rows"),
            ", ".join(str(row) for row in (issue.get("sample_rows") or [])),
            issue.get("message"),
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=column_index == 6, vertical="top")
    for index, width in enumerate([22, 24, 12, 14, 30, 70], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _write_header(sheet: Worksheet, headers: list[str], row: int = 1) -> None:
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _cell_value(finding: dict[str, Any], key: str) -> Any:
    """Flatten a finding field into something a cell can hold."""
    if key == "evidence_text":
        evidence = finding.get("evidence") or {}
        return "; ".join(f"{k}={v}" for k, v in evidence.items())
    value = finding.get(key)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def _disclaimer() -> str:
    """The disclaimer printed on every export."""
    return (
        "Findings are produced by deterministic Python rules from the uploaded file only. "
        "This report is not connected to a live SAP system, and no recommendation in it has been "
        "validated in an SAP environment. Any AI generated text is labelled as such and does not "
        "influence the risk determination."
    )
