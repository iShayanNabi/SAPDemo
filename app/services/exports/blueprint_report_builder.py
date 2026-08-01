"""Report builders for the SAP Blueprint Generator.

Four formats, all built from the same payload:

* **Markdown** - the format a blueprint actually travels in: it goes into a wiki,
  a pull request or a documentation site unchanged.
* **JSON** - the complete payload for a downstream system.
* **DOCX** - the format a project reviews in, with real headings so a table of
  contents can be generated and comments can be attached to a section.
* **PDF** - the read-only copy that gets forwarded.

Three things appear in **every** format, and they are the reason this module
exists rather than the exports being assembled in the API layer:

* the **disclaimer** - a blueprint is exactly the kind of document that gets
  forwarded away from the tool that produced it, so the statement that it is a
  proposal requiring review by qualified SAP professionals travels inside it;
* the **provenance of every section** - drafted by AI, computed from the project
  request, written from a template, or typed by a person;
* the **sections waiting for input** - a heading that says "nothing was written
  here, and here is the field to fill in" is far more useful in an exported
  document than a heading that has quietly been dropped.
"""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from app.core.logging import get_logger
from app.services.documents.pdf_writer import build_text_pdf

logger = get_logger(__name__)

#: Column order for an item table in the DOCX and Markdown exports.
ITEM_COLUMNS: tuple[tuple[str, str], ...] = (
    ("item_id", "ID"),
    ("title", "Item"),
    ("detail", "Detail"),
    ("category", "Category"),
    ("owner", "Owner"),
    ("rating", "Rating"),
    ("reference", "Reference"),
    ("source", "Written by"),
)

SOURCE_LABELS = {
    "ai_generated": "AI drafted",
    "mock_ai": "Mock AI drafted",
    "derived": "Computed from the project request",
    "template": "Written from the configured template",
    "manual": "Written by a person",
}

STATUS_LABELS = {
    "needs_input": "Waiting for project input",
    "draft": "Draft",
    "in_review": "In review",
    "approved": "Approved",
}


def blueprint_disclaimer(payload: dict[str, Any] | None = None) -> str:
    """The disclaimer printed on every blueprint export."""
    if payload and payload.get("disclaimer"):
        return str(payload["disclaimer"])
    return (
        "This is a PROPOSED SAP implementation blueprint drafted from a project request "
        "entered in this application. It requires review by qualified SAP professionals "
        "before any part of it is used to configure, build or plan a system. Nothing here has "
        "been validated against a live SAP system, and this application is not connected to "
        "one."
    )


def _source_label(section: dict[str, Any]) -> str:
    """Say plainly who wrote a section's words."""
    source = str(section.get("source", ""))
    if source == "ai_generated":
        origin = str(section.get("output_origin", ""))
        return SOURCE_LABELS.get(origin, SOURCE_LABELS["ai_generated"])
    return SOURCE_LABELS.get(source, source or "unknown")


def _approval_note(section: dict[str, Any]) -> str:
    """Render the approval and its staleness as one statement.

    Printing "approved by Ingrid" next to a section that describes a scope which
    changed after Ingrid read it is the single most misleading thing an export
    can do, so the two facts are never printed apart.
    """
    if not section.get("approved_by"):
        return ""
    note = f" · approved by {section['approved_by']}"
    if section.get("approval_is_stale"):
        stale = ", ".join(str(item) for item in section.get("stale_dependencies") or [])
        note += f" **before {stale} changed - re-review needed**"
    return note


def _status_label(section: dict[str, Any]) -> str:
    return STATUS_LABELS.get(str(section.get("status", "")), str(section.get("status", "")))


def _project_rows(project: dict[str, Any]) -> list[tuple[str, str]]:
    """The project request as label/value pairs, in a readable order."""
    order = (
        ("company", "Company"),
        ("industry", "Industry"),
        ("sap_product", "SAP product"),
        ("modules", "Modules"),
        ("business_objectives", "Business objectives"),
        ("countries", "Countries"),
        ("locations", "Locations"),
        ("company_codes", "Company codes"),
        ("plants", "Plants"),
        ("purchasing_organizations", "Purchasing organizations"),
        ("systems_involved", "Systems involved"),
        ("integrations", "Integrations"),
        ("data_sources", "Data sources"),
        ("user_groups", "User groups"),
        ("timeline", "Timeline"),
        ("constraints", "Constraints"),
        ("assumptions", "Assumptions"),
    )
    rows: list[tuple[str, str]] = []
    for key, label in order:
        value = project.get(key)
        if isinstance(value, list):
            text = "; ".join(str(item) for item in value) if value else "(not supplied)"
        else:
            text = str(value or "(not supplied)")
        rows.append((label, text))
    return rows


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------


def build_blueprint_json_report(payload: dict[str, Any]) -> bytes:
    """Serialise the full blueprint payload as JSON."""
    document = {
        "report_type": "sap_implementation_blueprint",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "disclaimer": blueprint_disclaimer(payload),
        **payload,
    }
    return json.dumps(document, indent=2, ensure_ascii=False, default=str).encode("utf-8")


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def build_blueprint_markdown_report(payload: dict[str, Any]) -> bytes:
    """Render the blueprint as Markdown."""
    blueprint = payload.get("blueprint", {}) or {}
    project = blueprint.get("project", {}) or {}
    summary = payload.get("summary", {}) or {}
    sections = payload.get("sections", []) or []

    lines: list[str] = [
        f"# {blueprint.get('name', 'SAP implementation blueprint')}",
        "",
        f"> **Proposed blueprint - requires review.** {blueprint_disclaimer(payload)}",
        "",
        f"*Generated at {datetime.now(timezone.utc).isoformat(timespec='seconds')} · "
        f"configuration v{blueprint.get('config_version', '')} · "
        f"engine v{blueprint.get('engine_version', '')} · "
        f"saved version {blueprint.get('current_version', 0)}*",
        "",
        "## Project request",
        "",
        "| Field | Value |",
        "| --- | --- |",
    ]
    for label, value in _project_rows(project):
        lines.append(f"| {label} | {_escape_cell(value)} |")

    lines += [
        "",
        "### Current process (as described by the business)",
        "",
        str(project.get("current_process") or "(not supplied)"),
        "",
        "### Desired process",
        "",
        str(project.get("desired_process") or "(not supplied)"),
        "",
        "## Document status",
        "",
        f"- Sections: **{summary.get('section_count', 0)}** "
        f"({summary.get('canonical_section_count', 0)} standard, "
        f"{summary.get('custom_section_count', 0)} custom)",
        f"- Written: **{summary.get('completeness_pct', 0)}%** · "
        f"Approved: **{summary.get('approval_pct', 0)}%** "
        f"({summary.get('approved_count', 0)} section(s))",
        f"- Waiting for project input: **{summary.get('needs_input_count', 0)}**",
        f"- Sections describing content that has since changed: "
        f"**{summary.get('stale_section_count', 0)}**"
        + (
            f" (of which **{summary.get('stale_approved_count', 0)}** still carry an "
            f"approval given before the change)"
            if summary.get("stale_approved_count")
            else ""
        ),
        f"- Items: **{summary.get('item_count', 0)}** · "
        f"AI drafted sections: {summary.get('ai_drafted_count', 0)} · "
        f"computed from the request: {summary.get('derived_count', 0)} · "
        f"from template: {summary.get('template_count', 0)} · "
        f"written by a person: {summary.get('manual_count', 0)}",
    ]
    if summary.get("missing_inputs"):
        missing = ", ".join(
            str(item).replace("_", " ") for item in summary["missing_inputs"]
        )
        lines.append(f"- Project fields still to supply: **{missing}**")
    if payload.get("ai_note"):
        lines += ["", f"*{payload['ai_note']}*"]

    lines += ["", "## Contents", ""]
    for section in sections:
        lines.append(
            f"{section.get('position', 0)}. **{section.get('title', '')}** "
            f"(`{section.get('section_id', '')}`) - {_status_label(section)}"
        )

    for section in sections:
        lines += ["", "---", "", f"## {section.get('position', 0)}. {section.get('title', '')}", ""]
        lines.append(
            f"`{section.get('section_id', '')}` · {_status_label(section)} · "
            f"{_source_label(section)}"
            + _approval_note(section)
        )
        if section.get("description"):
            lines += ["", f"*{section['description']}*"]
        if section.get("missing_inputs"):
            missing = ", ".join(
                str(item).replace("_", " ") for item in section["missing_inputs"]
            )
            lines += ["", f"> **Waiting for project input:** {missing}."]
        if section.get("stale_dependencies"):
            stale = ", ".join(str(item) for item in section["stale_dependencies"])
            lines += [
                "",
                f"> **Review needed:** this section was written before the following "
                f"section(s) changed: {stale}.",
            ]
        lines += ["", str(section.get("narrative") or "")]

        items = section.get("items") or []
        if items:
            lines += ["", *_item_markdown(section, items)]
        for note in section.get("validation_notes") or []:
            lines += ["", f"> Drafting note: {note}"]
        if section.get("comments"):
            lines += ["", f"> Reviewer comment: {section['comments']}"]

    versions = payload.get("versions") or []
    if versions:
        lines += ["", "---", "", "## Version history", "", "| Version | Label | Saved by | Sections | Approved |", "| --- | --- | --- | --- | --- |"]
        for version in versions:
            lines.append(
                f"| {version.get('version_number')} | {_escape_cell(version.get('label', ''))} "
                f"| {_escape_cell(version.get('created_by', ''))} "
                f"| {version.get('section_count', 0)} | {version.get('approved_count', 0)} |"
            )

    methodology = payload.get("methodology") or {}
    if methodology:
        lines += ["", "---", "", "## How this document was produced", ""]
        lines.append("**Decided by deterministic code:**")
        lines += [f"- {item}" for item in methodology.get("deterministic", [])]
        lines += ["", "**Written by a language model (when one was available):**"]
        lines += [f"- {item}" for item in methodology.get("ai_generated", [])]
        factual = methodology.get("factual_sections_ai_cannot_add_to") or []
        if factual:
            lines += [
                "",
                "**Sections a language model may describe but never add to:** "
                + ", ".join(str(item) for item in factual)
                + ".",
            ]
    if payload.get("review_note"):
        lines += ["", f"*{payload['review_note']}*"]

    lines += ["", "---", "", blueprint_disclaimer(payload), ""]
    return "\n".join(lines).encode("utf-8")


def _escape_cell(value: Any) -> str:
    """Make a value safe for a Markdown table cell."""
    return str(value or "").replace("|", "\\|").replace("\n", " ")


def _item_markdown(section: dict[str, Any], items: list[dict[str, Any]]) -> list[str]:
    """Render a section's items as a table or a list, per its content kind."""
    if str(section.get("content_kind")) == "table":
        columns = [
            (key, label)
            for key, label in ITEM_COLUMNS
            if key in ("item_id", "title", "detail", "category")
            or any(str(item.get(key) or "").strip() for item in items)
        ]
        lines = [
            "| " + " | ".join(label for _key, label in columns) + " |",
            "| " + " | ".join("---" for _ in columns) + " |",
        ]
        for item in items:
            lines.append(
                "| "
                + " | ".join(_escape_cell(_item_cell(item, key)) for key, _label in columns)
                + " |"
            )
        return lines

    lines = []
    for item in items:
        detail = f" - {item.get('detail')}" if item.get("detail") else ""
        suffix = f" *({item.get('category')})*" if item.get("category") else ""
        lines.append(f"- **{item.get('title', '')}**{suffix}{detail}")
    return lines


def _item_cell(item: dict[str, Any], key: str) -> str:
    if key == "source":
        return SOURCE_LABELS.get(str(item.get("source", "")), str(item.get("source", "")))
    return str(item.get(key) or "")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def build_blueprint_docx_report(payload: dict[str, Any]) -> bytes:
    """Render the blueprint as a Word document.

    Real heading styles are used rather than bold paragraphs, so Word can build
    a table of contents and a reviewer can comment on a section - which is the
    only reason to produce a DOCX rather than a PDF.
    """
    try:
        from docx import Document  # python-docx
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover - python-docx is a hard requirement
        raise RuntimeError(
            "python-docx is required for the DOCX export. Install it with: "
            "pip install -r requirements.txt"
        ) from exc

    blueprint = payload.get("blueprint", {}) or {}
    project = blueprint.get("project", {}) or {}
    summary = payload.get("summary", {}) or {}
    sections = payload.get("sections", []) or []

    document = Document()
    document.add_heading(str(blueprint.get("name", "SAP implementation blueprint")), level=0)

    warning = document.add_paragraph()
    run = warning.add_run("PROPOSED BLUEPRINT - REQUIRES REVIEW. ")
    run.bold = True
    warning.add_run(blueprint_disclaimer(payload))

    meta = document.add_paragraph()
    meta.add_run(
        f"Generated at {datetime.now(timezone.utc).isoformat(timespec='seconds')} · "
        f"configuration v{blueprint.get('config_version', '')} · "
        f"engine v{blueprint.get('engine_version', '')} · "
        f"saved version {blueprint.get('current_version', 0)}"
    ).italic = True

    document.add_heading("Project request", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in _project_rows(project):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("Current process (as described by the business)", level=2)
    document.add_paragraph(str(project.get("current_process") or "(not supplied)"))
    document.add_heading("Desired process", level=2)
    document.add_paragraph(str(project.get("desired_process") or "(not supplied)"))

    document.add_heading("Document status", level=1)
    for line in (
        f"Sections: {summary.get('section_count', 0)} "
        f"({summary.get('canonical_section_count', 0)} standard, "
        f"{summary.get('custom_section_count', 0)} custom)",
        f"Written: {summary.get('completeness_pct', 0)}% · "
        f"approved: {summary.get('approval_pct', 0)}% "
        f"({summary.get('approved_count', 0)} section(s))",
        f"Waiting for project input: {summary.get('needs_input_count', 0)}",
        f"Sections describing content that has since changed: "
        f"{summary.get('stale_section_count', 0)} "
        f"(approved despite the change: {summary.get('stale_approved_count', 0)})",
        f"Items: {summary.get('item_count', 0)}",
    ):
        document.add_paragraph(line, style="List Bullet")
    if payload.get("ai_note"):
        document.add_paragraph(str(payload["ai_note"])).runs[0].italic = True

    for section in sections:
        document.add_heading(
            f"{section.get('position', 0)}. {section.get('title', '')}", level=1
        )
        provenance = document.add_paragraph()
        stamp = provenance.add_run(
            f"{section.get('section_id', '')} · {_status_label(section)} · "
            f"{_source_label(section)}"
            + _approval_note(section)
        )
        stamp.italic = True
        stamp.font.size = Pt(9)

        if section.get("missing_inputs"):
            missing = ", ".join(
                str(item).replace("_", " ") for item in section["missing_inputs"]
            )
            document.add_paragraph(f"Waiting for project input: {missing}.").runs[0].bold = True
        if section.get("stale_dependencies"):
            stale = ", ".join(str(item) for item in section["stale_dependencies"])
            document.add_paragraph(
                f"Review needed: this section was written before the following section(s) "
                f"changed: {stale}."
            ).runs[0].bold = True

        for paragraph in str(section.get("narrative") or "").split("\n\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())

        items = section.get("items") or []
        if items and str(section.get("content_kind")) == "table":
            columns = [
                (key, label)
                for key, label in ITEM_COLUMNS
                if key in ("item_id", "title", "detail", "category")
                or any(str(item.get(key) or "").strip() for item in items)
            ]
            item_table = document.add_table(rows=1, cols=len(columns))
            item_table.style = "Light Grid Accent 1"
            for index, (_key, label) in enumerate(columns):
                item_table.rows[0].cells[index].text = label
            for item in items:
                cells = item_table.add_row().cells
                for index, (key, _label) in enumerate(columns):
                    cells[index].text = _item_cell(item, key)
        else:
            for item in items:
                text = str(item.get("title", ""))
                if item.get("detail"):
                    text += f" - {item['detail']}"
                document.add_paragraph(text, style="List Bullet")

        for note in section.get("validation_notes") or []:
            document.add_paragraph(f"Drafting note: {note}").runs[0].italic = True
        if section.get("comments"):
            document.add_paragraph(f"Reviewer comment: {section['comments']}")

    versions = payload.get("versions") or []
    if versions:
        document.add_heading("Version history", level=1)
        history = document.add_table(rows=1, cols=5)
        history.style = "Light Grid Accent 1"
        for index, label in enumerate(("Version", "Label", "Saved by", "Sections", "Approved")):
            history.rows[0].cells[index].text = label
        for version in versions:
            cells = history.add_row().cells
            cells[0].text = str(version.get("version_number", ""))
            cells[1].text = str(version.get("label", ""))
            cells[2].text = str(version.get("created_by", ""))
            cells[3].text = str(version.get("section_count", 0))
            cells[4].text = str(version.get("approved_count", 0))

    methodology = payload.get("methodology") or {}
    if methodology:
        document.add_heading("How this document was produced", level=1)
        document.add_paragraph("Decided by deterministic code:")
        for item in methodology.get("deterministic", []):
            document.add_paragraph(str(item), style="List Bullet")
        document.add_paragraph("Written by a language model (when one was available):")
        for item in methodology.get("ai_generated", []):
            document.add_paragraph(str(item), style="List Bullet")

    document.add_paragraph(blueprint_disclaimer(payload))

    buffer = io.BytesIO()
    document.save(buffer)
    logger.info("Built a DOCX blueprint with %d section(s)", len(sections))
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def build_blueprint_pdf_report(payload: dict[str, Any]) -> bytes:
    """Render the blueprint as a readable PDF.

    Built on the same dependency-free text PDF writer the Contract Assistant
    uses: one column, extractable text, no extra dependency for a format that
    only has to be read.
    """
    return build_text_pdf(
        _pdf_text(payload),
        title=str((payload.get("blueprint") or {}).get("name") or "SAP blueprint"),
    ).content


def _pdf_text(payload: dict[str, Any]) -> str:
    """Lay the blueprint out as plain text, one section per page."""
    blueprint = payload.get("blueprint", {}) or {}
    project = blueprint.get("project", {}) or {}
    summary = payload.get("summary", {}) or {}
    sections = payload.get("sections", []) or []
    lines: list[str] = []

    lines.append("SAP IMPLEMENTATION BLUEPRINT (PROPOSED - REQUIRES REVIEW)")
    lines.append("=" * 70)
    lines.append(f"Blueprint:    {blueprint.get('name', '')}")
    lines.append(f"Company:      {project.get('company', '')}")
    lines.append(f"Industry:     {project.get('industry', '')}")
    lines.append(f"SAP product:  {project.get('sap_product', '')}")
    lines.append(f"Generated at: {datetime.now(timezone.utc).isoformat(timespec='seconds')}")
    lines.append(f"Saved version: {blueprint.get('current_version', 0)}")
    lines.append("")
    lines.append(
        f"Sections: {summary.get('section_count', 0)}   "
        f"Written: {summary.get('completeness_pct', 0)}%   "
        f"Approved: {summary.get('approved_count', 0)}   "
        f"Waiting for input: {summary.get('needs_input_count', 0)}   "
        f"Needing re-review: {summary.get('stale_section_count', 0)}"
    )
    lines.append("")
    lines.append("Project request")
    lines.append("-" * 70)
    for label, value in _project_rows(project):
        lines.append(f"  {label + ':':<28}{value}")
    lines.append("")
    lines.append("Current process")
    lines.append("-" * 70)
    lines.append(str(project.get("current_process") or "(not supplied)"))
    lines.append("")
    lines.append("Desired process")
    lines.append("-" * 70)
    lines.append(str(project.get("desired_process") or "(not supplied)"))
    lines.append("")
    lines.append("Contents")
    lines.append("-" * 70)
    for section in sections:
        lines.append(
            f"  {str(section.get('position', 0)):>3}. {str(section.get('title', '')):<42}"
            f"{_status_label(section)}"
        )
    lines.append("")
    lines.append("Disclaimer")
    lines.append("-" * 70)
    lines.append(blueprint_disclaimer(payload))
    if payload.get("ai_note"):
        lines.append("")
        lines.append(str(payload["ai_note"]))

    for section in sections:
        lines.append("\f")
        lines.append(f"{section.get('position', 0)}. {section.get('title', '')}")
        lines.append("=" * 70)
        lines.append(f"Identifier:  {section.get('section_id', '')}")
        lines.append(f"Status:      {_status_label(section)}")
        lines.append(f"Written by:  {_source_label(section)}")
        lines.append(
            f"Approved by: {section.get('approved_by') or '-'}"
            + (
                "  (APPROVAL PREDATES A LATER CHANGE)"
                if section.get("approval_is_stale")
                else ""
            )
        )
        if section.get("missing_inputs"):
            missing = ", ".join(
                str(item).replace("_", " ") for item in section["missing_inputs"]
            )
            lines.append(f"WAITING FOR PROJECT INPUT: {missing}")
        if section.get("stale_dependencies"):
            stale = ", ".join(str(item) for item in section["stale_dependencies"])
            lines.append(
                f"REVIEW NEEDED: written before these section(s) changed: {stale}"
            )
        lines.append("")
        lines.append(str(section.get("narrative") or ""))

        items = section.get("items") or []
        if items:
            lines.append("")
            lines.append("-" * 70)
            for index, item in enumerate(items, start=1):
                lines.append(f"  {index}. {item.get('title', '')}")
                if item.get("category"):
                    lines.append(f"       Category: {item['category']}")
                if item.get("detail"):
                    lines.append(f"       {item['detail']}")
                for key, label in (("owner", "Owner"), ("rating", "Rating"), ("reference", "Ref")):
                    if item.get(key):
                        lines.append(f"       {label}: {item[key]}")
                lines.append(
                    f"       Written by: "
                    f"{SOURCE_LABELS.get(str(item.get('source', '')), item.get('source', ''))}"
                )
        for note in section.get("validation_notes") or []:
            lines.append(f"  Drafting note: {note}")
        if section.get("comments"):
            lines.append(f"  Reviewer comment: {section['comments']}")

    return "\n".join(lines)
