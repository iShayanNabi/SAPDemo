"""Text extractors for the three formats that need no external service.

* :class:`PdfTextExtractor` - text-based PDFs, via ``pypdf``.
* :class:`DocxExtractor` - Word documents, via ``python-docx``.
* :class:`PlainTextExtractor` - ``.txt`` / ``.md``.

None of them contacts a network service, so the first version of the Contract
Assistant works with no OCR provider at all. A PDF that parses but yields no
text is not treated as a failure: it comes back with ``needs_ocr=True`` and the
caller explains that the file looks scanned.
"""

from __future__ import annotations

import io
from typing import Any

from app.core.config import settings
from app.core.logging import get_logger
from app.services.documents.base import (
    DocumentExtractionError,
    DocumentExtractor,
    ExtractedPage,
    ExtractionResult,
    normalise_text,
)
from app.services.files.validation import decode_text

logger = get_logger(__name__)

#: DOCX and TXT have no page concept. When the document contains no explicit
#: page breaks the extractor falls back to a fixed character budget so a
#: citation still points at a stable, reproducible "page".
CHARS_PER_SYNTHETIC_PAGE = 3_000

#: The form feed a .txt author (or our own sample generator) uses for a break.
PAGE_BREAK = "\f"


class PdfTextExtractor(DocumentExtractor):
    """Extract per-page text from a text-based PDF."""

    name = "pdf_text"
    extensions = (".pdf",)

    def extract(self, content: bytes, *, filename: str = "") -> ExtractionResult:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - pypdf is a hard requirement
            raise DocumentExtractionError(
                "PDF support is not installed. Run: pip install -r requirements.txt"
            ) from exc

        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as exc:  # noqa: BLE001 - surfaced as a safe message
            logger.warning("PDF parsing failed: %s", type(exc).__name__)
            raise DocumentExtractionError(
                "The PDF could not be read. It may be corrupt or password protected."
            ) from exc

        if getattr(reader, "is_encrypted", False):
            # An empty user password is common for "protected" contracts and is
            # worth one attempt; anything else is reported plainly.
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                raise DocumentExtractionError(
                    "The PDF is password protected, so its text cannot be extracted. "
                    "Upload an unprotected copy."
                ) from None

        notes: list[str] = []
        pages: list[ExtractedPage] = []
        page_objects = list(reader.pages)
        if len(page_objects) > settings.max_document_pages:
            raise DocumentExtractionError(
                f"The PDF has {len(page_objects)} pages, which exceeds the "
                f"{settings.max_document_pages} page limit.",
                details={
                    "page_count": len(page_objects),
                    "max_pages": settings.max_document_pages,
                },
            )

        for index, page in enumerate(page_objects, start=1):
            try:
                raw = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001 - one bad page is not a bad document
                logger.warning("PDF page %d could not be read: %s", index, type(exc).__name__)
                notes.append(f"Page {index} could not be read and was treated as empty.")
                raw = ""
            pages.append(ExtractedPage(page_number=index, text=normalise_text(raw)))

        if not pages:
            raise DocumentExtractionError("The PDF contains no pages.")

        metadata: dict[str, Any] = {}
        try:
            info = reader.metadata or {}
            for key in ("/Title", "/Author", "/Subject", "/Producer"):
                value = info.get(key)
                if value:
                    metadata[key.lstrip("/").lower()] = str(value)[:200]
        except Exception:  # noqa: BLE001 - metadata is a nicety, never a failure
            pass

        result = ExtractionResult(
            pages=pages,
            extractor=self.name,
            source_format="pdf",
            page_basis="pdf_page",
            notes=notes,
            metadata=metadata,
        )
        _flag_if_scanned(result, "PDF")
        return result


class DocxExtractor(DocumentExtractor):
    """Extract text from a Word document, including its tables."""

    name = "docx"
    extensions = (".docx",)

    def extract(self, content: bytes, *, filename: str = "") -> ExtractionResult:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover - python-docx is a hard requirement
            raise DocumentExtractionError(
                "DOCX support is not installed. Run: pip install -r requirements.txt"
            ) from exc

        try:
            document = docx.Document(io.BytesIO(content))
        except Exception as exc:  # noqa: BLE001
            logger.warning("DOCX parsing failed: %s", type(exc).__name__)
            raise DocumentExtractionError(
                "The Word document could not be read. It may be corrupt or not a real .docx file."
            ) from exc

        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if _has_page_break(paragraph):
                blocks.append(PAGE_BREAK)
            if text:
                # A heading style is the most reliable structural signal a DOCX
                # carries; keeping it on its own line lets the section detector
                # find it exactly as it would in a PDF.
                blocks.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))

        body = "\n".join(blocks)
        notes: list[str] = []
        if document.tables:
            notes.append(
                f"{len(document.tables)} table(s) were flattened into pipe separated rows."
            )

        pages, basis = _split_into_pages(body)
        result = ExtractionResult(
            pages=pages,
            extractor=self.name,
            source_format="docx",
            page_basis=basis,
            notes=notes,
            metadata=_docx_metadata(document),
        )
        _flag_if_scanned(result, "Word document")
        return result


class PlainTextExtractor(DocumentExtractor):
    """Extract text from a plain text or Markdown file."""

    name = "plain_text"
    extensions = (".txt", ".md")

    def extract(self, content: bytes, *, filename: str = "") -> ExtractionResult:
        text = decode_text(content)
        if text is None:
            raise DocumentExtractionError(
                "The text file could not be decoded. Save it as UTF-8 and retry."
            )
        pages, basis = _split_into_pages(normalise_text(text))
        result = ExtractionResult(
            pages=pages,
            extractor=self.name,
            source_format="txt",
            page_basis=basis,
        )
        # A .txt cannot hold an image, so "short" never means "scanned" here.
        _flag_if_scanned(result, "text file", can_contain_images=False)
        return result


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _has_page_break(paragraph: Any) -> bool:
    """Return ``True`` when a DOCX paragraph starts a new page."""
    try:
        xml = paragraph._p.xml  # noqa: SLF001 - python-docx exposes no public API
    except Exception:  # noqa: BLE001
        return False
    return 'w:br w:type="page"' in xml or "lastRenderedPageBreak" in xml


def _docx_metadata(document: Any) -> dict[str, Any]:
    """Read the small amount of core metadata python-docx exposes."""
    metadata: dict[str, Any] = {}
    try:
        properties = document.core_properties
        for key in ("title", "author", "subject"):
            value = getattr(properties, key, None)
            if value:
                metadata[key] = str(value)[:200]
    except Exception:  # noqa: BLE001
        pass
    return metadata


def _split_into_pages(text: str) -> tuple[list[ExtractedPage], str]:
    """Split page-less text into pages, preferring real page breaks."""
    normalised = normalise_text(text)
    if PAGE_BREAK in normalised:
        chunks = [chunk.strip("\n") for chunk in normalised.split(PAGE_BREAK)]
        chunks = [chunk for chunk in chunks if chunk.strip()] or [""]
        return (
            [ExtractedPage(page_number=i, text=c) for i, c in enumerate(chunks, start=1)],
            "page_break",
        )

    if len(normalised) <= CHARS_PER_SYNTHETIC_PAGE:
        return [ExtractedPage(page_number=1, text=normalised)], "char_budget"

    # Break on paragraph boundaries so a clause is never cut in half by the
    # synthetic page boundary that it will later be cited against.
    chunks: list[str] = []
    current: list[str] = []
    length = 0
    for block in normalised.split("\n\n"):
        block_length = len(block) + 2
        if current and length + block_length > CHARS_PER_SYNTHETIC_PAGE:
            chunks.append("\n\n".join(current))
            current, length = [], 0
        current.append(block)
        length += block_length
    if current:
        chunks.append("\n\n".join(current))

    return (
        [ExtractedPage(page_number=i, text=c) for i, c in enumerate(chunks, start=1)],
        "char_budget",
    )


def _flag_if_scanned(
    result: ExtractionResult, label: str, *, can_contain_images: bool = True
) -> None:
    """Mark a result that parsed but produced (almost) no text.

    ``can_contain_images`` is what keeps the heuristic honest. A PDF or a DOCX
    page holding four characters is almost certainly a scan. A ``.txt`` file
    holding four characters is a short note: the format cannot carry an image,
    so the only conclusion available is "empty", never "scanned".
    """
    threshold = settings.min_chars_per_text_page if can_contain_images else 1
    text_pages = [page for page in result.pages if page.char_count >= threshold]
    if text_pages:
        if len(text_pages) < result.page_count:
            empty = result.page_count - len(text_pages)
            result.notes.append(
                f"{empty} of {result.page_count} page(s) held almost no text. Those pages are "
                f"most likely scanned images and were not analysed."
            )
        return

    result.needs_ocr = True
    if can_contain_images:
        result.notes.append(
            f"No selectable text was found in this {label}. It is most likely a scan or a set "
            f"of images. Configure an OCR provider, or upload a text-based version."
        )
    else:
        result.notes.append(f"This {label} is empty - it contains no text to analyse.")
